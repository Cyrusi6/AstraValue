from __future__ import annotations

import html
import json
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .event_state import current_event_versions
from .models import ReportVersion
from .registry import PROJECT_ROOT


DEFAULT_EXPORT_DIR = PROJECT_ROOT / "var" / "exports"
BLUE = "173B57"
LIGHT_BLUE = "E8F0F6"
WHITE = "FFFFFF"
TEXT = "17212B"
MUTED = "607080"
INPUT_BLUE = "0000FF"
LINK_GREEN = "008000"
THIN_GRAY = Side(style="thin", color="D6DEE5")


def _claim_detail(claim) -> str:
    parts = []
    if claim.evidence_ids:
        parts.append("原文/记录ID：" + ", ".join(claim.evidence_ids))
    for label, values in [("反证",claim.counter_evidence),("未知项",claim.unknowns),("失效条件",claim.invalidation_conditions)]:
        if values:
            parts.append(label + "：" + "；".join(values))
    return "；".join(parts)


def render_markdown(report: ReportVersion) -> str:
    c = report.conclusion
    used_methods = [run.method_ref for run in report.audit.model_runs if run.status.value == "成功"]
    lines = [
        f"# {report.company_name}（{report.ticker}）八步财报分析",
        "",
        "> 个人投研辅助材料，不构成投资建议或交易指令。",
        "",
        "## 置顶结论卡",
        "",
        f"- 报告状态：**{report.status.value}**",
        f"- 研究评级：**{c.rating.value}**（{'已由用户确认' if c.rating_confirmed else '系统建议，待用户确认'}）",
        f"- 公司 / 行业：{report.company_name}（{report.ticker}） / {report.industry}",
        f"- 数据截止 / 当前价格时点：{report.as_of.isoformat()} / {c.price_as_of.isoformat() if c.price_as_of else '暂无该数据'}",
        f"- 当前价格：{_display(c.current_price)}",
        f"- 合理价值区间：{_display(c.fair_value_low)} — {_display(c.fair_value_high)}（基准 {_display(c.fair_value_base)}）",
        f"- 安全边际：{_percent(c.margin_of_safety)}",
        f"- 已执行估值方法：{', '.join(used_methods) if used_methods else '暂不估值'}",
        f"- 证据完整度 / 可信度：{_percent(c.evidence_completeness)} / {_percent(c.evidence_confidence)}",
        f"- 八步数据覆盖快照：`{report.research_coverage_snapshot_id or '该版本未评估'}`",
        f"- 报告版本 / 数据快照 / 方法集合：v{report.version} / `{report.data_snapshot_id}` / `{report.method_bundle_id}`",
        "",
        "### 核心逻辑",
        *_bullets(c.core_theses),
        "",
        "### 主要风险",
        *_bullets(c.major_risks),
        "",
        "### 近期催化剂",
        *_bullets(c.catalysts),
        "",
        "### 结论失效条件",
        *_bullets(c.invalidation_conditions),
        "",
        "### 下一次跟踪事项",
        *_bullets(c.next_tracking_items),
        "",
    ]
    if report.request_metadata.get("report_notes"):
        lines.extend([f"> 备注：{report.request_metadata['report_notes']}", ""])

    lines.extend(
        [
            "## 八步数据覆盖",
            "",
            f"- 覆盖快照：`{report.research_coverage_snapshot_id or '该版本未评估'}`",
            f"- 数据状态：{report.research_coverage.get('overall_status', '未评估')}",
            f"- 研究窗口：{_stringify(report.research_coverage.get('analysis_scope', {})) or '未评估'}",
            "",
            _markdown_table(_coverage_rows(report)),
            "",
        ]
    )

    for section in report.sections:
        lines.extend([f"## {section.number}. {section.title}", "", section.summary, ""])
        if section.method_refs:
            lines.extend([f"方法引用：{', '.join(f'`{item}`' for item in section.method_refs)}", ""])
        if section.facts:
            lines.extend(["### 事实", "", *_bullets(section.facts), ""])
        if section.claims:
            lines.extend(["### 结论与证据", ""])
            for claim in section.claims:
                evidence = ", ".join(
                    [f"fact:{item}" for item in claim.evidence_fact_ids]
                    + [f"source:{item}" for item in claim.evidence_source_ids]
                ) or "分析判断"
                lines.append(f"- [{claim.claim_kind.value}] {claim.text}（证据：{evidence}）")
                if _claim_detail(claim):
                    lines.append(f"  {_claim_detail(claim)}")
            lines.append("")
        for table in section.tables:
            lines.extend([f"### {table['name']}", "", _markdown_table(table.get("rows", [])), ""])
        if section.warnings:
            lines.extend(["### 注意事项", "", *_bullets(section.warnings), ""])

    lines.extend(["## 审计附录", "", "### 来源清单", ""])
    source_rows = [
        {
            "source_id": item.source_id,
            "来源": item.name,
            "上游": item.upstream_source_id or "暂无该数据",
            "披露时间": item.published_at.isoformat() if item.published_at else "暂无该数据",
            "URL/位置": item.url or "本地归档",
            "哈希": item.document_hash or "暂无该数据",
        }
        for item in report.audit.sources
    ]
    lines.extend([_markdown_table(source_rows), "", "### 公式、方法与版本", ""])
    method_rows = [
        {
            "method_ref": item.ref,
            "状态": item.status,
            "文档": item.doc_path,
            "内容哈希": report.audit.method_bundle.method_hashes.get(item.ref, "暂无该数据"),
        }
        for item in report.audit.method_bundle.methods
    ]
    lines.extend([_markdown_table(method_rows), "", "### 情景假设", ""])
    assumption_rows = [
        {
            "assumption_id": item.assumption_id,
            "情景": item.scenario.value,
            "变量": item.name,
            "数值": item.value,
            "单位": item.unit,
            "来源": ",".join(item.source_ids) or "暂无该数据",
            "理由": item.reason,
            "有效期": item.valid_until.isoformat() if item.valid_until else "暂无该数据",
            "已确认": item.confirmed,
            "跟踪指标": item.tracking_metric or "暂无该数据",
        }
        for item in report.assumptions
    ]
    lines.extend([_markdown_table(assumption_rows), "", "### 模型运行与输入血缘", ""])
    model_rows = [
        {
            "run_id": item.run_id,
            "method_ref": item.method_ref,
            "状态": item.status.value,
            "输入血缘": item.input_lineage,
            "失败原因": item.failure_reason or "",
        }
        for item in report.audit.model_runs
    ]
    lines.extend([_markdown_table(model_rows), "", "### 双源核验记录", ""])
    lines.extend(
        [
            _markdown_table([item.model_dump(mode="json") for item in report.audit.verification_records]),
            "",
            "### 数据质量检查",
            "",
            _markdown_table(report.audit.data_quality_checks),
            "",
            "### 数据冲突",
            "",
            *_bullets(report.audit.conflicts),
            "",
            "### 缺失项",
            "",
            *_bullets(report.audit.missing_items),
            "",
            "### 无血缘数字",
            "",
            *_bullets(report.audit.unreferenced_numbers),
            "",
            "### 人工修订",
            "",
            *_bullets(report.audit.manual_edits),
            "",
            "### 版本差异",
            "",
            *_bullets(report.audit.version_changes),
            "",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def render_html(report: ReportVersion) -> str:
    c = report.conclusion
    status_class = "good" if report.status.value == "已复核" else "attention"
    sections = []
    for section in report.sections:
        claims_html = "".join(
            "<li><span class='badge'>{kind}</span>{text}<small>{evidence}</small></li>".format(
                kind=html.escape(claim.claim_kind.value),
                text=html.escape(claim.text),
                evidence=html.escape(
                    " / ".join(
                        [f"fact:{item}" for item in claim.evidence_fact_ids]
                        + [f"source:{item}" for item in claim.evidence_source_ids]
                        + [_claim_detail(claim)]
                    )
                    or "分析判断"
                ),
            )
            for claim in section.claims
        ) or "<li class='muted'>暂无经证据支持的结论</li>"
        facts_html = "".join(f"<li>{html.escape(item)}</li>" for item in section.facts)
        tables_html = "".join(
            f"<div class='table-block'><h3>{html.escape(table['name'])}</h3>{_html_table(table.get('rows', []))}</div>"
            for table in section.tables
        )
        warnings = "".join(f"<li>{html.escape(item)}</li>" for item in section.warnings)
        sections.append(
            f"<section class='step'><div class='step-no'>0{section.number}</div><div class='step-content'>"
            f"<h2>{html.escape(section.title)}</h2><p class='lede'>{html.escape(section.summary)}</p>"
            f"<p class='method-ref'>{html.escape(' · '.join(section.method_refs))}</p>"
            f"<ul class='claims'>{claims_html}{facts_html}</ul>{tables_html}"
            f"{f'<aside class=warning><strong>注意事项</strong><ul>{warnings}</ul></aside>' if warnings else ''}"
            "</div></section>"
        )

    source_rows = [
        {
            "ID": item.source_id,
            "来源": item.name,
            "上游": item.upstream_source_id,
            "披露时间": item.published_at,
            "位置": item.url or "本地归档",
        }
        for item in report.audit.sources
    ]
    method_rows = [
        {
            "方法版本": item.ref,
            "状态": item.status,
            "文档": item.doc_path,
            "哈希": report.audit.method_bundle.method_hashes.get(item.ref),
        }
        for item in report.audit.method_bundle.methods
    ]
    assumptions = [
        {
            "情景": item.scenario.value,
            "变量": item.name,
            "值": item.value,
            "单位": item.unit,
            "理由": item.reason,
            "来源": item.source_ids,
            "确认": item.confirmed,
        }
        for item in report.assumptions
    ]
    coverage_rows = _coverage_rows(report)
    coverage_print_rows = _coverage_print_rows(report)
    audit_blocks = "".join(
        [
            "<h3>来源清单</h3>" + _html_table(source_rows),
            "<h3>公式、方法与版本</h3>" + _html_table(method_rows),
            "<h3>情景假设</h3>" + _html_table(assumptions),
            "<h3>模型运行与输入血缘</h3>" + _html_table([item.model_dump(mode="json") for item in report.audit.model_runs]),
            "<h3>双源核验记录</h3>" + _html_table([item.model_dump(mode="json") for item in report.audit.verification_records]),
            "<h3>数据质量检查</h3>" + _html_table(report.audit.data_quality_checks),
            "<h3>冲突、缺失与人工修订</h3>"
            + _html_table(
                [{"类型": "数据冲突", "内容": item} for item in report.audit.conflicts]
                + [{"类型": "缺失项", "内容": item} for item in report.audit.missing_items]
                + [{"类型": "无血缘数字", "内容": item} for item in report.audit.unreferenced_numbers]
                + [{"类型": "人工修订", "内容": item} for item in report.audit.manual_edits]
                + [{"类型": "版本变化", "内容": item} for item in report.audit.version_changes]
            ),
        ]
    )
    note = report.request_metadata.get("report_notes")
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(report.company_name)}八步财报分析</title><style>{_CSS}</style></head>
<body><main>
<header class="hero"><div><p class="eyebrow">A股全行业八步财报分析 · v{report.version}</p>
<h1>{html.escape(report.company_name)} <span>{html.escape(report.ticker)}</span></h1>
<p>{html.escape(report.industry)} · 数据截止 {report.as_of.isoformat()}</p></div>
<div class="status {status_class}">{html.escape(report.status.value)}</div></header>
<section class="conclusion"><div class="section-heading"><p>RESEARCH VIEW</p><h2>置顶结论</h2></div>
<div class="kpis">
<article><label>研究评级</label><strong>{html.escape(c.rating.value)}</strong><small>{'用户已确认' if c.rating_confirmed else '系统建议，待确认'}</small></article>
<article><label>当前价格</label><strong>{_display(c.current_price)}</strong><small>{c.price_as_of.isoformat() if c.price_as_of else '暂无时点'}</small></article>
<article><label>合理价值区间</label><strong>{_display(c.fair_value_low)} - {_display(c.fair_value_high)}</strong><small>基准 {_display(c.fair_value_base)}</small></article>
<article><label>安全边际</label><strong>{_percent(c.margin_of_safety)}</strong><small>相对基准价值</small></article>
</div>
<div class="triad"><div><h3>核心逻辑</h3>{_html_list(c.core_theses)}</div><div><h3>主要风险</h3>{_html_list(c.major_risks)}</div><div><h3>失效条件</h3>{_html_list(c.invalidation_conditions)}</div></div>
<p class="meta">证据完整度 {_percent(c.evidence_completeness)} · 可信度 {_percent(c.evidence_confidence)} · 八步数据覆盖 {html.escape(report.research_coverage_snapshot_id or '该版本未评估')} · 快照 {html.escape(report.data_snapshot_id)} · 方法 {html.escape(report.method_bundle_id)}</p>
{f'<aside class="fixture-note">{html.escape(str(note))}</aside>' if note else ''}</section>
<section class="audit"><div class="section-heading"><p>DATA READINESS</p><h2>八步数据覆盖</h2></div>
<p class="meta">覆盖快照 {html.escape(report.research_coverage_snapshot_id or '该版本未评估')} · 数据状态 {html.escape(str(report.research_coverage.get('overall_status', '未评估')))}</p>
<div class="screen-only">{_html_table(coverage_rows)}</div>
<div class="print-only">{_html_table(coverage_print_rows, table_class='coverage-print')}</div></section>
{''.join(sections)}
<section class="audit"><div class="section-heading"><p>AUDIT TRAIL</p><h2>审计附录</h2></div>{audit_blocks}</section>
<footer>个人投研辅助材料，不构成投资建议。评级需由用户确认，所有数字应通过审计附录复核。</footer>
</main></body></html>"""


def export_report(report: ReportVersion, fmt: str, export_dir: Path | str = DEFAULT_EXPORT_DIR) -> Path:
    export_dir = Path(export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_ticker = "".join(char for char in report.ticker if char.isalnum() or char in "-_")
    base = export_dir / f"{safe_ticker}_v{report.version}_{report.report_id[:8]}"
    fmt = fmt.lower()
    if fmt in {"md", "markdown"}:
        path = base.with_suffix(".md")
        path.write_text(render_markdown(report), encoding="utf-8")
        return path
    if fmt == "html":
        path = base.with_suffix(".html")
        path.write_text(render_html(report), encoding="utf-8")
        return path
    if fmt in {"xlsx", "excel"}:
        path = base.with_suffix(".xlsx")
        _write_xlsx(report, path)
        return path
    if fmt == "pdf":
        html_path = base.with_suffix(".html")
        html_path.write_text(render_html(report), encoding="utf-8")
        pdf_path = base.with_suffix(".pdf")
        _html_to_pdf(html_path, pdf_path)
        return pdf_path
    raise ValueError(f"不支持的导出格式: {fmt}")


def _write_xlsx(report: ReportVersion, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    summary = workbook.create_sheet("置顶结论")
    sections_sheet = workbook.create_sheet("八步正文")
    facts_sheet = workbook.create_sheet("财务事实")
    dimensional_facts_sheet = workbook.create_sheet("经营维度事实")
    events_sheet = workbook.create_sheet("公司事件")
    assumptions_sheet = workbook.create_sheet("情景假设")
    scenarios_sheet = workbook.create_sheet("情景测算")
    valuation_sheet = workbook.create_sheet("估值模型")
    audit_sheet = workbook.create_sheet("来源与审计")
    methods_sheet = workbook.create_sheet("方法与版本")
    coverage_sheet = workbook.create_sheet("八步数据覆盖")
    checks_sheet = workbook.create_sheet("检查")
    _write_summary(summary, report)
    _write_sections(sections_sheet, report)
    if report.request_metadata.get("display_metrics"):
        table_sheet=workbook.create_sheet("核心期间表")
        table_rows=[]
        for table in report.sections[1].tables:
            for row in table.get("rows",[]):
                for key,value in row.items():
                    if key not in {"指标","口径"}:
                        table_rows.append([row.get("指标"),row.get("口径"),key,value])
        _append_table(table_sheet,["指标","口径","期间","显示值与事实引用"],table_rows,[32,16,16,46])
        index_sheet=workbook.create_sheet("核心事实索引")
        _append_table(index_sheet,["短引用","fact_id","期间类型","单位"],
            [[m.get("fact_ref"),m["fact"]["fact_id"],m["fact"]["period_type"],m["fact"]["unit"]]
             for m in report.request_metadata["display_metrics"] if m.get("state")=="ready"],[12,65,20,16])
    _write_facts(facts_sheet, report)
    _write_dimensional_facts(dimensional_facts_sheet, report)
    _write_events(events_sheet, report)
    assumption_rows = _write_assumptions(assumptions_sheet, report)
    _write_scenarios(scenarios_sheet, report, assumption_rows)
    _write_valuations(valuation_sheet, report)
    _write_audit(audit_sheet, report)
    _write_methods(methods_sheet, report)
    _write_coverage(coverage_sheet, report)
    _write_checks(checks_sheet, report)
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.35
        sheet.page_margins.right = 0.35
        sheet.page_margins.top = 0.55
        sheet.page_margins.bottom = 0.55
    workbook.save(path)


def _write_summary(sheet, report: ReportVersion) -> None:
    c = report.conclusion
    sheet.merge_cells("A1:H2")
    sheet["A1"] = f"{report.company_name}（{report.ticker}）八步财报分析"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=22, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 27
    metadata = [
        ("报告状态", report.status.value), ("研究评级", c.rating.value),
        ("评级确认", "已确认" if c.rating_confirmed else "待用户确认"), ("行业", report.industry),
        ("报告版本", report.version), ("数据截止", report.as_of.isoformat()),
        ("数据快照", report.data_snapshot_id), ("方法集合", report.method_bundle_id),
        ("八步覆盖快照", report.research_coverage_snapshot_id or "该版本未评估"),
        ("八步数据状态", report.research_coverage.get("overall_status", "未评估")),
    ]
    for index, (label, value) in enumerate(metadata):
        row = 4 + index // 2
        col = 1 + (index % 2) * 4
        sheet.cell(row, col, label).font = Font(bold=True, color=MUTED)
        sheet.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 3)
        sheet.cell(row, col + 1, value)
    kpis = [
        ("当前价格", c.current_price, "0.00"), ("价值下限", c.fair_value_low, "0.00"),
        ("基准价值", c.fair_value_base, "0.00"), ("价值上限", c.fair_value_high, "0.00"),
        ("安全边际", c.margin_of_safety, "0.0%"), ("证据完整度", c.evidence_completeness, "0.0%"),
        ("证据可信度", c.evidence_confidence, "0.0%"),
    ]
    for index, (label, value, fmt) in enumerate(kpis):
        col = index + 1
        sheet.cell(9, col, label).font = Font(bold=True, color=MUTED)
        sheet.cell(10, col, value)
        sheet.cell(10, col).number_format = fmt
        sheet.cell(10, col).font = Font(size=15, bold=True, color=TEXT)
        sheet.cell(10, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    row = 13
    for title, items in (("核心逻辑", c.core_theses), ("主要风险", c.major_risks), ("失效条件", c.invalidation_conditions)):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        sheet.cell(row, 1, title)
        _style_section_header(sheet.cell(row, 1))
        row += 1
        for item in items or ["暂无该数据"]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            sheet.cell(row, 1, f"• {item}").alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=8)
    sheet.cell(row, 1, report.request_metadata.get("report_notes") or "个人投研辅助材料，不构成投资建议。")
    sheet.cell(row, 1).fill = PatternFill("solid", fgColor="FFF4D6")
    sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    _set_widths(sheet, [17] * 8)
    sheet.freeze_panes = "A9"


def _write_sections(sheet, report: ReportVersion) -> None:
    if report.request_metadata.get("display_metrics"):
        rows=[]
        for section in report.sections:
            for claim in section.claims:
                rows.append([section.number,section.title,claim.text,
                    ", ".join(claim.evidence_fact_ids+claim.evidence_source_ids),_claim_detail(claim)])
            if not section.claims:
                rows.append([section.number,section.title,section.summary,"","；".join(section.warnings)])
        _append_table(sheet,["步骤","标题","研究观察","事实与来源","原文定位、反证和边界"],rows,[8,26,80,60,70])
        import math
        for row_index,row in enumerate(rows,2):
            lines=max(math.ceil(sum(2 if ord(c)>127 else 1 for c in str(v))/width) for v,width in zip(row,[8,26,80,60,70]))
            sheet.row_dimensions[row_index].height=min(409,18*lines+12)
        return
    headers = ["步骤", "标题", "摘要", "结论及证据", "方法引用", "注意事项"]
    rows = []
    for section in report.sections:
        claims = "\n".join(
            f"[{item.claim_kind.value}] {item.text} | fact:{','.join(item.evidence_fact_ids)} source:{','.join(item.evidence_source_ids)}"
            for item in section.claims
        )
        rows.append([section.number, section.title, section.summary, claims or "暂无经证据支持的结论", "\n".join(section.method_refs), "\n".join(section.warnings)])
    _append_table(sheet, headers, rows, [9, 25, 42, 64, 34, 48])


def _write_facts(sheet, report: ReportVersion) -> None:
    headers = ["fact_id", "metric_id", "value", "unit", "period_start", "period_end", "period_type", "disclosed_at", "as_of", "scope", "audited", "status", "source_ids", "method_ref", "derived_from", "page", "restatement"]
    rows = []
    source_map = {item.source_id: item for item in report.audit.sources}
    for fact in report.facts:
        rows.append([
            fact.fact_id, fact.metric_id, fact.value, fact.unit,
            fact.period_start.isoformat() if fact.period_start else None,
            fact.period_end.isoformat() if fact.period_end else None, fact.period_type,
            fact.disclosed_at.isoformat() if fact.disclosed_at else None, fact.as_of.isoformat(), fact.scope,
            fact.audited, fact.verification_status.value, ",".join(fact.source_ids), fact.method_ref,
            ",".join(fact.derived_from_fact_ids), fact.document_page, fact.restatement_version,
        ])
    _append_table(sheet, headers, rows, [36, 25, 18, 14, 13, 13, 16, 24, 24, 16, 10, 14, 32, 28, 32, 8, 16])
    for row_index, fact in enumerate(report.facts, 2):
        urls = [source_map[item].url for item in fact.source_ids if item in source_map and source_map[item].url]
        if urls:
            sheet.cell(row_index, 3).comment = Comment("来源：\n" + "\n".join(urls), "User")
        if fact.value is not None:
            sheet.cell(row_index, 3).number_format = ('0.00"倍"' if fact.metric_id in {"eastmoney_pe_ttm","eastmoney_pb_mrq","eastmoney_ps_ttm"} else _excel_number_format(fact.unit))
    sheet.freeze_panes = "C2"


def _write_dimensional_facts(sheet, report: ReportVersion) -> None:
    headers = [
        "dimensional_fact_id", "ticker", "metric_id", "dimension_type",
        "dimension_name", "dimension_code", "parent_dimension", "value",
        "unit", "currency", "period_start", "period_end", "period_type",
        "available_at", "scope", "accounting_basis", "audited", "status",
        "source_ids", "document_ids", "page", "table_title", "row_label",
        "column_label", "evidence_text", "evidence_offsets",
        "extraction_method", "supporting_fact_ids", "method_ref",
        "data_snapshot_id", "disclosed_as", "metadata",
    ]
    rows = []
    source_map = {item.source_id: item for item in report.audit.sources}
    for fact in report.dimensional_facts:
        spans = fact.evidence_spans
        document_ids = _join_unique(
            [*fact.document_ids, *(item.document_id for item in spans)]
        )
        pages = _join_unique(
            [fact.document_page, *(item.page for item in spans)]
        )
        table_titles = _join_unique(
            [fact.table_title, *(item.table_title for item in spans)]
        )
        row_labels = _join_unique(
            [fact.row_label, *(item.row_label for item in spans)]
        )
        column_labels = _join_unique(
            [fact.column_label, *(item.column_label for item in spans)]
        )
        evidence_text = "\n---\n".join(
            item.text.strip() for item in spans if item.text and item.text.strip()
        )
        evidence_offsets = _join_unique(
            f"{item.start_offset if item.start_offset is not None else ''}:"
            f"{item.end_offset if item.end_offset is not None else ''}"
            for item in spans
            if item.start_offset is not None or item.end_offset is not None
        )
        rows.append([
            fact.dimensional_fact_id,
            fact.ticker,
            fact.metric_id,
            fact.dimension_type,
            fact.dimension_name,
            fact.dimension_code,
            fact.parent_dimension,
            fact.value,
            fact.unit,
            fact.currency,
            fact.period_start.isoformat() if fact.period_start else None,
            fact.period_end.isoformat() if fact.period_end else None,
            fact.period_type,
            fact.available_at.isoformat(),
            fact.scope,
            fact.accounting_basis,
            fact.audited,
            fact.verification_status.value,
            ",".join(fact.source_ids),
            document_ids,
            pages,
            table_titles,
            row_labels,
            column_labels,
            evidence_text or None,
            evidence_offsets or None,
            fact.extraction_method.value,
            ",".join(fact.supporting_fact_ids),
            fact.method_ref,
            fact.data_snapshot_id,
            fact.disclosed_as.value,
            _stringify(fact.metadata) if fact.metadata else None,
        ])
    _append_table(
        sheet,
        headers,
        rows,
        [
            38, 13, 26, 18, 28, 18, 24, 18, 13, 10, 13, 13, 15, 24,
            16, 22, 10, 14, 36, 38, 10, 34, 28, 24, 64, 18, 20, 36,
            28, 28, 14, 54,
        ],
    )
    for row_index, fact in enumerate(report.dimensional_facts, 2):
        urls = [
            source_map[item].url
            for item in fact.source_ids
            if item in source_map and source_map[item].url
        ]
        if urls:
            sheet.cell(row_index, 8).comment = Comment(
                "来源：\n" + "\n".join(urls),
                "User",
            )
        if fact.value is not None:
            sheet.cell(row_index, 8).number_format = _excel_number_format(fact.unit)
    sheet.freeze_panes = "H2"


def _write_events(sheet, report: ReportVersion) -> None:
    headers = [
        "event_id", "ticker", "event_type", "event_subtype",
        "lifecycle_state", "is_current", "announced_at", "available_at",
        "effective_at", "period_start", "period_end", "summary", "amount",
        "currency", "shares", "ratio", "materiality", "status",
        "root_event_id", "previous_event_id", "source_ids", "document_ids",
        "page", "evidence_text", "evidence_offsets", "parties", "event_terms",
        "affected_metrics", "supporting_fact_ids", "claim_kind",
        "status_updated_at", "data_snapshot_id", "metadata",
    ]
    current_ids = {
        item.event_id for item in current_event_versions(report.events)
    }
    source_map = {item.source_id: item for item in report.audit.sources}
    rows = []
    for event in report.events:
        pages = _join_unique(
            item.page for item in event.evidence_spans if item.page is not None
        )
        evidence_text = "\n---\n".join(
            item.text.strip()
            for item in event.evidence_spans
            if item.text and item.text.strip()
        )
        offsets = _join_unique(
            f"{item.start_offset if item.start_offset is not None else ''}:"
            f"{item.end_offset if item.end_offset is not None else ''}"
            for item in event.evidence_spans
            if item.start_offset is not None or item.end_offset is not None
        )
        rows.append([
            event.event_id,
            event.ticker,
            event.event_type,
            event.event_subtype,
            event.lifecycle_state,
            event.event_id in current_ids,
            event.announced_at.isoformat(),
            event.available_at.isoformat(),
            event.effective_at.isoformat() if event.effective_at else None,
            event.period_start.isoformat() if event.period_start else None,
            event.period_end.isoformat() if event.period_end else None,
            event.summary,
            event.amount,
            event.currency,
            event.shares,
            event.ratio,
            event.materiality,
            event.verification_status.value,
            event.root_event_id,
            event.previous_event_id,
            ",".join(event.source_ids),
            ",".join(event.document_ids),
            pages,
            evidence_text or None,
            offsets or None,
            _stringify(event.parties) if event.parties else None,
            _stringify(event.event_terms) if event.event_terms else None,
            ",".join(event.affected_metrics),
            ",".join(event.supporting_fact_ids),
            event.claim_kind.value,
            event.status_updated_at.isoformat(),
            event.data_snapshot_id,
            _stringify(event.metadata) if event.metadata else None,
        ])
    _append_table(
        sheet,
        headers,
        rows,
        [
            38, 13, 22, 22, 18, 12, 24, 24, 24, 13, 13, 54, 18, 12,
            18, 14, 14, 14, 38, 38, 36, 38, 10, 64, 18, 42, 54, 30,
            36, 16, 24, 28, 54,
        ],
    )
    for row_index, event in enumerate(report.events, 2):
        urls = [
            source_map[item].url
            for item in event.source_ids
            if item in source_map and source_map[item].url
        ]
        if urls:
            sheet.cell(row_index, 12).comment = Comment(
                "来源：\n" + "\n".join(urls),
                "User",
            )
        if event.amount is not None:
            sheet.cell(row_index, 13).number_format = _excel_number_format(
                event.currency
            )
        if event.shares is not None:
            sheet.cell(row_index, 15).number_format = _excel_number_format("share")
        if event.ratio is not None:
            sheet.cell(row_index, 16).number_format = _excel_number_format("ratio")
    sheet.freeze_panes = "G2"


def _write_assumptions(sheet, report: ReportVersion) -> dict[tuple[str, str], int]:
    headers = ["情景", "变量", "数值", "单位", "理由", "来源ID", "有效期", "已确认", "跟踪指标", "assumption_id"]
    rows = []
    row_map = {}
    for index, item in enumerate(report.assumptions, 2):
        row_map[(item.scenario.value, item.name)] = index
        rows.append([item.scenario.value, item.name, item.value, item.unit, item.reason, ",".join(item.source_ids), item.valid_until.isoformat() if item.valid_until else None, item.confirmed, item.tracking_metric, item.assumption_id])
    _append_table(sheet, headers, rows, [10, 29, 18, 13, 40, 28, 13, 10, 24, 38])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 3).font = Font(color=INPUT_BLUE)
        sheet.cell(row, 3).fill = PatternFill("solid", fgColor="FFF9D9")
        sheet.cell(row, 3).number_format = _excel_number_format(sheet.cell(row, 4).value)
    sheet.freeze_panes = "C2"
    return row_map


def _write_scenarios(sheet, report: ReportVersion, row_map: dict[tuple[str, str], int]) -> None:
    result_rows = report.sections[7].tables[0]["rows"] if report.sections[7].tables else []
    result_map = {item["scenario"]: item for item in result_rows}
    row = 1
    required = {
        "base_revenue", "volume_growth", "price_growth", "gross_margin", "expense_ratio", "tax_rate",
        "depreciation_ratio", "capex_ratio", "working_capital_ratio", "wacc", "terminal_growth",
        "years", "net_debt", "shares_outstanding", "minority_interest_ratio", "non_recurring_after_tax",
    }
    for scenario in ("悲观", "基准", "乐观"):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        sheet.cell(row, 1, f"{scenario}情景 - 公式驱动测算")
        _style_section_header(sheet.cell(row, 1))
        growth_row = row + 1
        sheet.cell(growth_row, 1, "收入增速")
        missing = sorted(name for name in required if (scenario, name) not in row_map)
        if missing:
            sheet.cell(growth_row, 2, "暂无该数据")
            sheet.merge_cells(
                start_row=growth_row,
                start_column=3,
                end_row=growth_row,
                end_column=9,
            )
            sheet.cell(growth_row, 3, "缺失假设：" + "、".join(missing))
            sheet.cell(growth_row, 3).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            sheet.row_dimensions[growth_row].height = 42
            row = growth_row + 3
            continue
        sheet.cell(growth_row, 2, f"=(1+{_assumption_ref(row_map, scenario, 'volume_growth')})*(1+{_assumption_ref(row_map, scenario, 'price_growth')})-1")
        sheet.cell(growth_row, 2).number_format = "0.0%"
        sheet.cell(growth_row, 2).font = Font(color=LINK_GREEN)
        header_row = row + 3
        for col, value in enumerate(["年度", "营收", "EBIT", "净利润", "扣非归母", "EPS", "FCFF", "FCFF现值", "说明"], 1):
            sheet.cell(header_row, col, value)
            _style_header(sheet.cell(header_row, col))
        years = int(_assumption_value(report, scenario, "years") or 0)
        first_projection = header_row + 1
        for year in range(1, years + 1):
            current = first_projection + year - 1
            sheet.cell(current, 1, year)
            previous_revenue = _assumption_ref(row_map, scenario, "base_revenue") if year == 1 else f"B{current - 1}"
            sheet.cell(current, 2, f"={previous_revenue}*(1+$B${growth_row})")
            gm = _assumption_ref(row_map, scenario, "gross_margin")
            expense = _assumption_ref(row_map, scenario, "expense_ratio")
            tax = _assumption_ref(row_map, scenario, "tax_rate")
            minority = _assumption_ref(row_map, scenario, "minority_interest_ratio")
            nonrec = _assumption_ref(row_map, scenario, "non_recurring_after_tax")
            shares = _assumption_ref(row_map, scenario, "shares_outstanding")
            depreciation = _assumption_ref(row_map, scenario, "depreciation_ratio")
            capex = _assumption_ref(row_map, scenario, "capex_ratio")
            wc = _assumption_ref(row_map, scenario, "working_capital_ratio")
            wacc = _assumption_ref(row_map, scenario, "wacc")
            sheet.cell(current, 3, f"=B{current}*({gm}-{expense})")
            sheet.cell(current, 4, f"=C{current}*(1-{tax})")
            sheet.cell(current, 5, f"=D{current}*(1-{minority})-{nonrec}")
            sheet.cell(current, 6, f"=E{current}/{shares}")
            sheet.cell(current, 7, f"=D{current}+B{current}*{depreciation}-B{current}*{capex}-(B{current}-{previous_revenue})*{wc}")
            sheet.cell(current, 8, f"=G{current}/(1+{wacc})^A{current}")
            sheet.cell(current, 9, "由情景假设跨表驱动")
            for col in range(2, 9):
                sheet.cell(current, col).font = Font(color=LINK_GREEN)
                sheet.cell(current, col).number_format = "0.00" if col == 6 else "#,##0;[Red](#,##0);-"
        summary_row = first_projection + years + 1
        if years:
            last = first_projection + years - 1
            wacc = _assumption_ref(row_map, scenario, "wacc")
            terminal_growth = _assumption_ref(row_map, scenario, "terminal_growth")
            net_debt = _assumption_ref(row_map, scenario, "net_debt")
            shares = _assumption_ref(row_map, scenario, "shares_outstanding")
            sheet.cell(summary_row, 1, "终值")
            sheet.cell(summary_row, 2, f"=G{last}*(1+{terminal_growth})/({wacc}-{terminal_growth})")
            sheet.cell(summary_row + 1, 1, "公式合理价值/股")
            sheet.cell(summary_row + 1, 2, f"=(SUM(H{first_projection}:H{last})+B{summary_row}/(1+{wacc})^{years}-{net_debt})/{shares}")
            sheet.cell(summary_row + 2, 1, "报告冻结值/股")
            sheet.cell(summary_row + 2, 2, result_map.get(scenario, {}).get("fair_value_per_share"))
            sheet.cell(summary_row + 3, 1, "公式-冻结值")
            sheet.cell(summary_row + 3, 2, f"=B{summary_row + 1}-B{summary_row + 2}")
            for current in range(summary_row, summary_row + 4):
                sheet.cell(current, 1).font = Font(bold=True)
                sheet.cell(current, 2).number_format = "0.0000"
            sheet.cell(summary_row + 1, 2).font = Font(color=LINK_GREEN, bold=True)
        row = summary_row + 6
    _set_widths(sheet, [18, 19, 19, 19, 19, 14, 19, 19, 30])
    sheet.freeze_panes = "A4"


def _write_valuations(sheet, report: ReportVersion) -> None:
    row = 1
    for run in report.audit.model_runs:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.cell(row, 1, f"{run.method_ref} · {run.status.value}")
        _style_section_header(sheet.cell(row, 1))
        row += 1
        for col, value in enumerate(["输入项", "数值", "输入血缘", "冻结输出", "说明"], 1):
            sheet.cell(row, col, value)
            _style_header(sheet.cell(row, col))
        input_start = row + 1
        key_rows = {}
        for key, value in run.inputs.items():
            row += 1
            key_rows[key] = row
            sheet.cell(row, 1, key)
            sheet.cell(row, 2, _cell_value(value))
            sheet.cell(row, 3, _stringify(run.input_lineage.get(key, [])))
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                sheet.cell(row, 2).font = Font(color=INPUT_BLUE)
        output = run.outputs or {}
        if row >= input_start:
            sheet.cell(input_start, 4, output.get("fair_value_per_share") if isinstance(output, dict) else None)
            sheet.cell(input_start, 4).number_format = "0.0000"
        row += 1
        sheet.cell(row, 1, "公式复核")
        formula = _valuation_formula(run.method_ref, key_rows)
        if formula:
            sheet.cell(row, 2, formula)
            sheet.cell(row, 2).number_format = "0.0000"
            sheet.cell(row, 5, "可编辑输入，公式用于复核；冻结结果来自报告JSON")
        else:
            sheet.cell(row, 2, "见情景测算或方法配置")
            sheet.cell(row, 5, run.failure_reason or "复杂模型保留输入、输出与方法版本")
        row += 3
    _set_widths(sheet, [29, 31, 65, 18, 52])


def _write_audit(sheet, report: ReportVersion) -> None:
    headers = ["source_id", "来源", "类型", "上游来源", "URL/位置", "披露时间", "获取时间", "文档哈希", "权威等级", "备注"]
    rows = [[item.source_id, item.name, item.source_type, item.upstream_source_id, item.url or "本地归档", item.published_at.isoformat() if item.published_at else None, item.retrieved_at.isoformat(), item.document_hash, item.authority_level, item.notes] for item in report.audit.sources]
    _append_table(sheet, headers, rows, [30, 36, 24, 30, 55, 24, 24, 40, 10, 42])
    row = sheet.max_row + 3
    for title, items in (("数据冲突", report.audit.conflicts), ("缺失项", report.audit.missing_items), ("无血缘数字", report.audit.unreferenced_numbers), ("人工修订", report.audit.manual_edits), ("版本变化", report.audit.version_changes)):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        sheet.cell(row, 1, title)
        _style_section_header(sheet.cell(row, 1))
        row += 1
        for item in items or ["无"]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            sheet.cell(row, 1, item)
            row += 1
        row += 1


def _write_methods(sheet, report: ReportVersion) -> None:
    headers = ["method_ref", "状态", "适用行业", "文档", "实现", "测试", "内容哈希", "局限"]
    rows = [[item.ref, item.status, ",".join(item.applies_to), item.doc_path, item.implementation_ref, item.test_ref, report.audit.method_bundle.method_hashes.get(item.ref), "；".join(item.limitations)] for item in report.audit.method_bundle.methods]
    _append_table(sheet, headers, rows, [28, 10, 28, 48, 38, 42, 66, 44])


def _write_coverage(sheet, report: ReportVersion) -> None:
    rows = _coverage_rows(report)
    headers = list(rows[0])
    _append_table(
        sheet,
        headers,
        [[row.get(header) for header in headers] for row in rows],
        [36, 12, 18, 14, 14, 34, 34, 34, 24, 34, 48, 42, 16, 16, 16],
    )


def _coverage_rows(report: ReportVersion) -> list[dict[str, Any]]:
    """Project the frozen coverage payload identically into every export.

    Coverage is stored as JSON so old reports remain readable after the
    executable requirement registry changes.  Accept the current question
    projection and the smaller step-only development projection, but never
    infer readiness from the legacy evidence-completeness score.
    """

    coverage = report.research_coverage or {}
    snapshot_id = report.research_coverage_snapshot_id or "该版本未评估"
    analysis_scope = _stringify(coverage.get("analysis_scope", {})) or "未评估"
    questions = coverage.get("questions") or coverage.get("question_coverage") or []
    rows: list[dict[str, Any]] = []
    for question in questions:
        if not isinstance(question, dict):
            continue
        state = question.get("state", question.get("status", "pending"))
        applicability = question.get("applicability", "unknown")
        ready = question.get("ready_requirement_ids", question.get("ready_required_ids", []))
        missing = question.get("missing_requirement_ids", question.get("missing_required_ids", []))
        not_applicable = question.get("not_applicable_requirement_ids", [])
        required = question.get("required_requirement_ids")
        if required is None:
            required = _unique_values(ready, missing, not_applicable)
        optional_summary = {
            "ready": question.get("optional_ready_ids", []),
            "missing": question.get("optional_missing_ids", []),
        }
        if state == "not_applicable" and not not_applicable:
            not_applicable = ["问题整体不适用"]
        rows.append(
            {
                "覆盖快照ID": snapshot_id,
                "步骤": question.get("step_id", "未登记"),
                "问题": question.get("question_id", "未登记"),
                "数据状态": state,
                "适用性": applicability,
                "必需项": _stringify(required),
                "已得必需项": _stringify(ready),
                "待补必需项": _stringify(missing),
                "不适用项": _stringify(not_applicable),
                "可选增强": _stringify(optional_summary),
                "研究窗口": _stringify(question.get("analysis_scope", analysis_scope)),
                "补充路径": _stringify(
                    question.get("next_paths", question.get("supplement_paths", []))
                ),
                "方法状态": question.get("method_status", "unknown"),
                "假设状态": question.get("assumption_status", "unknown"),
                "研究状态": question.get("analysis_status", "not_started"),
            }
        )
    if rows:
        return rows

    steps = coverage.get("steps") or coverage.get("step_coverage") or []
    for step in steps:
        if not isinstance(step, dict):
            continue
        rows.append(
            {
                "覆盖快照ID": snapshot_id,
                "步骤": step.get("step_id", "未登记"),
                "问题": "步骤汇总",
                "数据状态": step.get("state", step.get("status", "pending")),
                "适用性": "见逐题明细",
                "必需项": "见逐题明细",
                "已得必需项": _stringify(step.get("ready_question_ids", [])),
                "待补必需项": _stringify(step.get("pending_question_ids", [])),
                "不适用项": _stringify(step.get("not_applicable_question_ids", [])),
                "可选增强": "见逐题明细",
                "研究窗口": analysis_scope,
                "补充路径": "见逐题明细",
                "方法状态": step.get("method_status", "unknown"),
                "假设状态": step.get("assumption_status", "unknown"),
                "研究状态": step.get("analysis_status", "not_started"),
            }
        )
    if rows:
        return rows

    return [
        {
            "覆盖快照ID": snapshot_id,
            "步骤": "ES01-ES08",
            "问题": "未评估",
            "数据状态": "未评估",
            "适用性": "unknown",
            "必需项": "未评估",
            "已得必需项": "未评估",
            "待补必需项": "未评估",
            "不适用项": "未评估",
            "可选增强": "未评估",
            "研究窗口": analysis_scope,
            "补充路径": "未评估",
            "方法状态": "unknown",
            "假设状态": "unknown",
            "研究状态": "not_started",
        }
    ]


def _coverage_print_rows(report: ReportVersion) -> list[dict[str, Any]]:
    """Keep PDF coverage readable while retaining the complete HTML table."""

    rows = []
    for row in _coverage_rows(report):
        required_count = _json_list_count(row.get("必需项"))
        ready_count = _json_list_count(row.get("已得必需项"))
        progress = (
            f"{ready_count}/{required_count}"
            if required_count is not None and ready_count is not None
            else "见完整明细"
        )
        rows.append(
            {
                "步骤": row.get("步骤"),
                "问题": row.get("问题"),
                "数据状态": row.get("数据状态"),
                "适用性": row.get("适用性"),
                "要求进度": progress,
                "待补必需项": row.get("待补必需项"),
                "补充路径": row.get("补充路径"),
                "方法/假设/研究": " / ".join(
                    str(row.get(key, "unknown"))
                    for key in ("方法状态", "假设状态", "研究状态")
                ),
            }
        )
    return rows


def _json_list_count(value: Any) -> int | None:
    if isinstance(value, (list, tuple, set)):
        return len(value)
    if not isinstance(value, str):
        return None
    try:
        parsed = json.loads(value)
    except (TypeError, ValueError):
        return None
    return len(parsed) if isinstance(parsed, list) else None


def _unique_values(*groups: Any) -> list[Any]:
    values: list[Any] = []
    for group in groups:
        if isinstance(group, (list, tuple, set)):
            candidates = group
        elif group in (None, ""):
            candidates = ()
        else:
            candidates = (group,)
        for item in candidates:
            if item not in values:
                values.append(item)
    return values


def _write_checks(sheet, report: ReportVersion) -> None:
    rows = [[item.get("check"), item.get("actual_difference"), 0, item.get("actual_difference"), item.get("tolerance"), item.get("status"), item.get("message"), item.get("method_ref")] for item in report.audit.data_quality_checks]
    rows.extend([
        ["严格八步", len(report.sections), 8, len(report.sections) - 8, 0, "OK" if len(report.sections) == 8 else "FAIL", "章节必须为1至8", "ReportVersion schema"],
        ["无血缘估值输入", len(report.audit.unreferenced_numbers), 0, len(report.audit.unreferenced_numbers), 0, "OK" if not report.audit.unreferenced_numbers else "FAIL", "所有模型数值输入必须有fact/assumption/source/formula引用", "ReportBuilder"],
        ["未解决冲突", len(report.audit.conflicts), 0, len(report.audit.conflicts), 0, "OK" if not report.audit.conflicts else "FAIL", "冲突输入不得进入确定性估值", "source_policy.json"],
        ["评级确认", int(report.conclusion.rating_confirmed), 1, int(report.conclusion.rating_confirmed) - 1, 0, "OK" if report.conclusion.rating_confirmed else "WARN", "草稿允许未确认，已复核报告必须确认", "report_policies.json"],
    ])
    _append_table(
        sheet,
        ["检查", "实际差异", "期望差异", "偏差", "容差", "状态", "说明", "规则/方法"],
        rows,
        [28, 16, 16, 16, 16, 12, 52, 32],
    )
    status_range = f"F2:F{sheet.max_row}"
    sheet.conditional_formatting.add(status_range, FormulaRule(formula=["F2=\"OK\""], fill=PatternFill("solid", fgColor="DDF2E2")))
    sheet.conditional_formatting.add(status_range, FormulaRule(formula=["F2=\"FAIL\""], fill=PatternFill("solid", fgColor="FADBD8")))
    sheet.conditional_formatting.add(status_range, FormulaRule(formula=["F2=\"WARN\""], fill=PatternFill("solid", fgColor="FFF1C2")))


def _append_table(sheet, headers: list[str], rows: Iterable[list[Any]], widths: list[int]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        _style_header(cell)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
            if cell.font.color is None or cell.font.color.type == "theme":
                cell.font = Font(name="Microsoft YaHei", size=10, color=TEXT)
    _set_widths(sheet, widths)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 26


def _style_header(cell) -> None:
    cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def _style_section_header(cell) -> None:
    cell.font = Font(name="Microsoft YaHei", bold=True, color=WHITE, size=12)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(vertical="center")


def _set_widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 80)


def _join_unique(values: Iterable[Any]) -> str:
    return "\n".join(
        dict.fromkeys(
            str(value)
            for value in values
            if value is not None and str(value).strip()
        )
    )


def _assumption_ref(row_map: dict[tuple[str, str], int], scenario: str, name: str) -> str:
    row = row_map.get((scenario, name))
    return f"'情景假设'!$C${row}" if row else "#N/A"


def _assumption_value(report: ReportVersion, scenario: str, name: str) -> float | None:
    return next((item.value for item in report.assumptions if item.scenario.value == scenario and item.name == name), None)


def _valuation_formula(method_ref: str, rows: dict[str, int]) -> str | None:
    def cell(key: str) -> str | None:
        return f"B{rows[key]}" if key in rows else None
    method = method_ref.split("@", 1)[0]
    if method == "VAL.RELATIVE" and cell("metric_per_share") and cell("target_multiple"):
        return f"={cell('metric_per_share')}*{cell('target_multiple')}"
    if method == "VAL.CYCLE_NORMALIZED" and cell("normalized_eps") and cell("target_pe"):
        return f"={cell('normalized_eps')}*{cell('target_pe')}"
    if method == "VAL.PEV" and cell("embedded_value_per_share") and cell("target_pev"):
        return f"={cell('embedded_value_per_share')}*{cell('target_pev')}"
    return None


def _html_to_pdf(html_path: Path, pdf_path: Path) -> None:
    candidates = [Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"), Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"), Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")]
    browser = next((item for item in candidates if item.exists()), None)
    if browser is None:
        raise RuntimeError("未找到Edge或Chrome，无法从统一HTML生成PDF")
    with tempfile.TemporaryDirectory(
        prefix=".ashare-pdf-",
        dir=pdf_path.parent.resolve(),
        ignore_cleanup_errors=True,
    ) as profile:
        generated_pdf = Path(profile) / "rendered.pdf"
        command = [str(browser), "--headless=new", "--disable-gpu", "--disable-background-networking", "--disable-component-update", "--no-first-run", "--host-resolver-rules=MAP * ~NOTFOUND", "--no-pdf-header-footer", "--print-to-pdf-no-header", "--run-all-compositor-stages-before-draw", "--virtual-time-budget=10000", f"--user-data-dir={profile}", f"--print-to-pdf={generated_pdf}", html_path.resolve().as_uri()]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=90, creationflags=0x08000000)
        # Edge sometimes hands PDF finalization to a child process and returns
        # before the file is complete. Keep the temporary profile alive briefly
        # and replace the destination only after the new artifact can be opened.
        # A unique source path also prevents a stale destination from passing QA.
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            if _promote_generated_pdf(generated_pdf, pdf_path):
                return
            time.sleep(0.2)
    detail = completed.stderr.strip() or completed.stdout.strip() or f"浏览器退出码 {completed.returncode}"
    raise RuntimeError(f"PDF生成失败: {detail}")


def _is_valid_pdf(path: Path) -> bool:
    if not path.exists() or path.stat().st_size < 1000:
        return False
    try:
        import fitz

        with fitz.open(path) as document:
            return document.is_pdf and document.page_count > 0
    except (ImportError, OSError, RuntimeError, ValueError):
        return False


def _promote_generated_pdf(generated_pdf: Path, pdf_path: Path) -> bool:
    """Publish a validated browser PDF even while Edge briefly locks its source."""
    if not _is_valid_pdf(generated_pdf):
        return False
    try:
        generated_pdf.replace(pdf_path)
        return _is_valid_pdf(pdf_path)
    except OSError:
        # Edge can keep the generated file open after the headless parent exits.
        # Windows still permits a read-only copy, so validate a sibling staging
        # file before atomically publishing it as the requested destination.
        staging = pdf_path.with_suffix(f"{pdf_path.suffix}.tmp")
        try:
            shutil.copy2(generated_pdf, staging)
            if not _is_valid_pdf(staging):
                return False
            staging.replace(pdf_path)
            return _is_valid_pdf(pdf_path)
        except OSError:
            return False
        finally:
            try:
                staging.unlink(missing_ok=True)
            except OSError:
                pass


def _normalize_rows(rows: Any) -> list[dict[str, Any]]:
    if isinstance(rows, dict):
        return [{"项目": key, "值": value} for key, value in rows.items()]
    if not isinstance(rows, list):
        return [{"值": rows}]
    normalized = [item if isinstance(item, dict) else {"值": item} for item in rows]
    if any(
        {"event_id", "event_type", "event_subtype", "状态", "摘要"} <= set(item)
        for item in normalized
    ):
        return _event_display_rows(normalized)
    if any({"scenario", "projections"} <= set(item) for item in normalized):
        return _scenario_display_rows(normalized)
    if any({"run_id", "method_ref", "input_lineage"} <= set(item) for item in normalized):
        return _model_run_display_rows(normalized)
    if any({"method_ref", "fair_value_per_share", "sensitivity"} <= set(item) for item in normalized):
        return _valuation_display_rows(normalized)
    return normalized


def _event_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    displayed: list[dict[str, Any]] = []
    for item in rows:
        numeric_parts = []
        if item.get("金额") is not None:
            numeric_parts.append(
                f"金额={_stringify(item.get('金额'))} {item.get('币种') or ''}".strip()
            )
        if item.get("股数") is not None:
            numeric_parts.append(f"股数={_stringify(item.get('股数'))}")
        if item.get("比例") is not None:
            ratio = item.get("比例")
            ratio_text = f"{ratio:.4%}" if isinstance(ratio, (int, float)) else _stringify(ratio)
            numeric_parts.append(f"比例={ratio_text}")

        state_parts = [
            str(value)
            for value in (item.get("event_subtype"), item.get("状态"))
            if value not in (None, "")
        ]
        if item.get("当前状态"):
            state_parts.append(f"当前={item.get('当前状态')}")

        chain_parts = []
        if item.get("root_event_id"):
            chain_parts.append(f"root:{item.get('root_event_id')}")
        if item.get("previous_event_id"):
            chain_parts.append(f"prev:{item.get('previous_event_id')}")

        evidence_parts = []
        source_ids = item.get("source_ids")
        if source_ids:
            evidence_parts.append(f"source:{_stringify(source_ids)}")
        document_ids = item.get("document_ids")
        if document_ids:
            evidence_parts.append(f"document:{_stringify(document_ids)}")
        pages = item.get("页码")
        if pages:
            evidence_parts.append(f"页:{_stringify(pages)}")

        announced_at = item.get("公告时间")
        displayed.append(
            {
                "事件ID": item.get("event_id"),
                "类别": item.get("事件类别") or item.get("event_type"),
                "状态/当前": " / ".join(state_parts) or "暂无该数据",
                "公告日期": str(announced_at)[:10] if announced_at else "暂无该数据",
                "摘要": item.get("摘要"),
                "关键数值": "；".join(numeric_parts) or "暂无该数据",
                "核验": item.get("核验状态"),
                "状态链": "；".join(chain_parts) or "暂无该数据",
                "证据定位": "；".join(evidence_parts) or item.get("lineage"),
            }
        )
    return displayed


def _scenario_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    displayed: list[dict[str, Any]] = []
    for item in rows:
        projections = item.get("projections")
        if isinstance(projections, list) and projections:
            for projection in projections:
                displayed.append(
                    {
                        "情景": item.get("scenario"),
                        "年度": projection.get("year"),
                        "收入增速": item.get("revenue_growth"),
                        "营收": projection.get("revenue"),
                        "扣非归母": projection.get("net_income_excl_parent"),
                        "EPS": projection.get("eps"),
                        "FCF": projection.get("fcff"),
                        "合理价值/股": item.get("fair_value_per_share"),
                        "状态/确认": "已确认" if item.get("confirmed") else "待人工确认",
                    }
                )
            continue
        detail = item.get("failure_reason")
        if not detail and item.get("missing"):
            detail = f"缺少: {', '.join(str(value) for value in item['missing'])}"
        displayed.append(
            {
                "情景": item.get("scenario"),
                "年度": None,
                "收入增速": None,
                "营收": None,
                "扣非归母": None,
                "EPS": None,
                "FCF": None,
                "合理价值/股": None,
                "状态/确认": f"{item.get('status', '待核验')}：{detail or '暂无该数据'}",
            }
        )
    return displayed


def _model_run_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    displayed: list[dict[str, Any]] = []
    output_fields = ("fair_value_per_share", "range_low", "range_high", "currency", "warnings")
    for item in rows:
        method_ref = item.get("method_ref")
        status = item.get("status")
        displayed.append(
            {
                "方法": method_ref,
                "状态": status,
                "类型": "运行信息",
                "项目": "run_id",
                "值": item.get("run_id"),
                "血缘/说明": item.get("created_at"),
            }
        )
        inputs = item.get("inputs") if isinstance(item.get("inputs"), dict) else {}
        lineage = item.get("input_lineage") if isinstance(item.get("input_lineage"), dict) else {}
        for key, value in inputs.items():
            displayed.append(
                {
                    "方法": method_ref,
                    "状态": status,
                    "类型": "输入",
                    "项目": key,
                    "值": value,
                    "血缘/说明": lineage.get(key) or "暂无该数据",
                }
            )
        outputs = item.get("outputs") if isinstance(item.get("outputs"), dict) else {}
        for key in output_fields:
            if key in outputs and outputs[key] not in (None, [], ""):
                displayed.append(
                    {
                        "方法": method_ref,
                        "状态": status,
                        "类型": "输出",
                        "项目": key,
                        "值": outputs[key],
                        "血缘/说明": f"{method_ref}确定性计算",
                    }
                )
        if item.get("failure_reason"):
            displayed.append(
                {
                    "方法": method_ref,
                    "状态": status,
                    "类型": "失败原因",
                    "项目": "failure_reason",
                    "值": item["failure_reason"],
                    "血缘/说明": "模型停止条件",
                }
            )
    return displayed


def _valuation_display_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "方法": item.get("method_ref"),
            "状态": item.get("status"),
            "合理价值/股": item.get("fair_value_per_share"),
            "区间下限": item.get("range_low"),
            "区间上限": item.get("range_high"),
            "币种": item.get("currency"),
            "提示": item.get("warnings") or "暂无该数据",
            "失败原因": item.get("failure_reason") or "暂无该数据",
        }
        for item in rows
    ]


def _markdown_table(rows: Any) -> str:
    normalized = _normalize_rows(rows)
    if not normalized:
        return "暂无该数据"
    columns = []
    for row in normalized:
        for key in row:
            if key not in columns:
                columns.append(key)
    header = "| " + " | ".join(str(item) for item in columns) + " |"
    divider = "| " + " | ".join("---" for _ in columns) + " |"
    body = ["| " + " | ".join(_markdown_cell(row.get(key)) for key in columns) + " |" for row in normalized]
    return "\n".join([header, divider, *body])


def _markdown_cell(value: Any) -> str:
    return _stringify(value).replace("|", "\\|").replace("\n", "<br>")


def _html_table(rows: Any, *, table_class: str = "") -> str:
    normalized = _normalize_rows(rows)
    if not normalized:
        return "<p class='muted'>暂无该数据</p>"
    columns = []
    for row in normalized:
        for key in row:
            if key not in columns:
                columns.append(key)
    head = "".join(f"<th>{html.escape(str(item))}</th>" for item in columns)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(_stringify(row.get(key)))}</td>" for key in columns) + "</tr>" for row in normalized)
    class_attribute = f" class='{html.escape(table_class)}'" if table_class else ""
    return f"<div class='table-scroll'><table{class_attribute}><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>"


def _stringify(value: Any) -> str:
    if value is None or value == "":
        return "暂无该数据"
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))
    return str(value)


def _cell_value(value: Any) -> Any:
    return value if isinstance(value, (str, int, float, bool)) or value is None else _stringify(value)


def _html_list(items: list[str]) -> str:
    return "<ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in (items or ["暂无该数据"])) + "</ul>"


def _bullets(items: list[str]) -> list[str]:
    return [f"- {item}" for item in items] if items else ["- 暂无该数据"]


def _display(value: Any) -> str:
    return "暂无该数据" if value is None else f"{value:.4f}" if isinstance(value, float) else str(value)


def _percent(value: float | None) -> str:
    return "暂无该数据" if value is None else f"{value:.2%}"


def _excel_number_format(unit: Any) -> str:
    text = str(unit or "")
    if text in {"ratio", "%"}:
        return "0.0%;[Red](0.0%);-"
    if "share" in text or "股" in text:
        return "#,##0.00;[Red](#,##0.00);-"
    if text in {"x", "multiple"}:
        return "0.0x;[Red](0.0x);-"
    return "#,##0;[Red](#,##0);-"


_CSS = """
:root{--ink:#17212b;--muted:#657483;--navy:#173b57;--blue:#2e678f;--pale:#f2f6f8;--line:#d9e1e7;--green:#1d7a45;--amber:#a56600;--red:#a6302b}*{box-sizing:border-box}body{margin:0;background:#e9eef1;color:var(--ink);font:14px/1.65 Inter,"Microsoft YaHei","Noto Sans CJK SC",sans-serif}main{max-width:1180px;margin:0 auto;padding:34px 28px 70px}.hero{display:flex;justify-content:space-between;gap:20px;align-items:flex-end;background:var(--navy);color:white;padding:32px 36px;border-radius:18px 18px 0 0}.hero h1{font-size:34px;line-height:1.2;margin:5px 0}.hero h1 span{font-size:18px;font-weight:500;opacity:.75}.hero p{margin:0}.eyebrow{letter-spacing:.12em;font-size:12px;color:#b9d6e9}.status{padding:7px 14px;border-radius:999px;font-weight:700;background:#fff;color:var(--navy)}.status.good{color:var(--green)}.status.attention{color:var(--amber)}section{background:white;border:1px solid var(--line);border-top:0;padding:30px 36px}.conclusion{border-top:4px solid #d5a940}.section-heading p{margin:0;color:var(--blue);font-size:11px;font-weight:800;letter-spacing:.16em}.section-heading h2{font-size:26px;margin:2px 0 20px}.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}.kpis article{padding:17px;border:1px solid var(--line);background:var(--pale);border-radius:9px}.kpis label,.kpis small{display:block;color:var(--muted)}.kpis strong{display:block;font-size:21px;margin:3px 0}.triad{display:grid;grid-template-columns:repeat(3,1fr);gap:24px;margin-top:24px}.triad h3{font-size:15px;border-bottom:2px solid var(--navy);padding-bottom:7px}.meta,.method-ref,small,.muted{color:var(--muted);font-size:12px}.fixture-note,.warning{margin-top:20px;padding:12px 15px;border-left:4px solid #d5a940;background:#fff8df}.step{display:grid;grid-template-columns:62px 1fr;gap:20px}.step-no{font:700 23px/1 Georgia,serif;color:var(--blue);padding-top:4px}.step h2{margin:0;font-size:23px}.lede{font-size:15px}.claims{padding-left:18px}.claims li{margin:7px 0}.claims small{display:block}.badge{display:inline-block;margin-right:7px;padding:1px 8px;border-radius:99px;background:#e2eef5;color:var(--blue);font-size:11px}.table-block{margin:20px 0}.table-block h3,.audit h3{font-size:15px;margin:22px 0 8px}.table-scroll{overflow:auto;border:1px solid var(--line);border-radius:7px}table{width:100%;border-collapse:collapse;font-size:12px}th{background:var(--navy);color:white;text-align:left;position:sticky;top:0}th,td{padding:8px 10px;border-bottom:1px solid var(--line);vertical-align:top;white-space:pre-wrap;word-break:normal;overflow-wrap:anywhere}tr:last-child td{border-bottom:0}.audit{border-top:5px solid var(--navy)}.print-only{display:none}footer{text-align:center;background:var(--navy);color:#c9d7e1;padding:22px;border-radius:0 0 18px 18px;font-size:12px}@media(max-width:800px){main{padding:0}.hero{border-radius:0}.kpis,.triad{grid-template-columns:1fr 1fr}.step{grid-template-columns:42px 1fr}section{padding:24px 18px}}@media print{@page{size:A4;margin:13mm 11mm}body{background:white;font-size:10px;print-color-adjust:exact;-webkit-print-color-adjust:exact}main{max-width:none;margin:0;padding:0}.hero{border-radius:0;padding:20px 24px}.hero h1{font-size:24px}section{padding:18px 22px}.step{break-before:page;display:block}.step-no{float:right}.step h2{font-size:19px}.kpis{grid-template-columns:repeat(4,1fr)}.kpis article{padding:10px}.kpis strong{font-size:15px}.triad{gap:12px}.table-scroll{overflow:visible}table{font-size:8px;table-layout:fixed}th,td{padding:4px;word-break:normal;overflow-wrap:anywhere}.screen-only{display:none}.print-only{display:block}.coverage-print th:nth-child(1){width:7%}.coverage-print th:nth-child(2){width:9%}.coverage-print th:nth-child(3){width:8%}.coverage-print th:nth-child(4){width:8%}.coverage-print th:nth-child(5){width:8%}.coverage-print th:nth-child(6){width:30%}.coverage-print th:nth-child(7){width:20%}.coverage-print th:nth-child(8){width:10%}.audit{break-before:page}.audit h3{break-after:avoid}footer{border-radius:0}.status{font-size:10px}}
"""
