from __future__ import annotations

import html
from dataclasses import asdict, dataclass
from datetime import datetime
from io import BytesIO
from typing import Callable
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.xml.functions import tostring

from .canonical import canonical_datetime, canonical_sha256
from .codex_tools import canonical_session_manifest_hash
from .models import CodexSessionManifest, GovernanceReport
from .redaction import redact_text
from .report_service import canonical_governance_report_hash


REPORT_VIEW_VERSION = "1.0.0"
MARKDOWN_RENDERER_VERSION = "1.0.0"
HTML_RENDERER_VERSION = "1.0.0"
XLSX_RENDERER_VERSION = "1.0.1"
PDF_RENDERER_VERSION = "1.0.0"

_FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_FIXED_WORKBOOK_TIMESTAMP = datetime(1980, 1, 1)


class ReportRenderError(ValueError):
    """Raised when report identities disagree or an output cannot be rendered safely."""


@dataclass(frozen=True, slots=True)
class GovernanceFindingView:
    finding_id: str
    layer: str
    text: str
    citation_ids: tuple[str, ...]
    uncertainty: str | None
    source_role: str | None


@dataclass(frozen=True, slots=True)
class GovernanceSectionView:
    section_id: str
    title: str
    findings: tuple[GovernanceFindingView, ...]


@dataclass(frozen=True, slots=True)
class GovernanceValidationCheckView:
    check_code: str
    passed: bool
    detail: str | None


@dataclass(frozen=True, slots=True)
class GovernanceReportView:
    """The sole immutable projection consumed by every report renderer."""

    schema_version: str
    renderer_version: str
    report_id: str
    report_hash: str
    company_id: str
    state_at: datetime
    known_at: datetime
    perspective: str
    snapshot_id: str
    snapshot_hash: str
    evidence_manifest_id: str
    evidence_manifest_hash: str
    session_manifest_id: str
    session_manifest_hash: str
    generation_status: str
    decision_author: str
    future_knowledge_used: bool
    sections: tuple[GovernanceSectionView, ...]
    citation_ids: tuple[str, ...]
    active_gap_ids: tuple[str, ...]
    active_conflict_ids: tuple[str, ...]
    pending_candidate_ids: tuple[str, ...]
    data_limitations: tuple[str, ...]
    technical_validation_passed: bool
    hard_failure_codes: tuple[str, ...]
    technical_checks: tuple[GovernanceValidationCheckView, ...]
    failure_code: str | None
    session_parent_report_run_id: str
    session_model_profile: str
    session_runner_protocol_version: str
    session_tool_protocol_version: str
    session_input_pack_id: str
    session_input_pack_hash: str
    session_initial_snapshot_id: str
    session_final_snapshot_id: str
    session_tool_read_ids: tuple[str, ...]
    session_research_task_ids: tuple[str, ...]
    session_research_result_bundle_ids: tuple[str, ...]
    session_snapshot_adoption_ids: tuple[str, ...]
    session_final_citation_ids: tuple[str, ...]

    @property
    def canonical_view_hash(self) -> str:
        return canonical_sha256(
            asdict(self),
            schema_name="governance-report-view",
            schema_version=self.renderer_version,
        )


def build_governance_report_view(
    report: GovernanceReport,
    session_manifest: CodexSessionManifest,
    *,
    renderer_version: str = REPORT_VIEW_VERSION,
) -> GovernanceReportView:
    """Bind a report and its session once before any format-specific rendering."""

    if not renderer_version.strip():
        raise ReportRenderError("renderer_version must be non-empty")
    if report.canonical_report_hash != canonical_governance_report_hash(report):
        raise ReportRenderError("report canonical hash does not match report content")
    if session_manifest.canonical_hash != canonical_session_manifest_hash(
        session_manifest
    ):
        raise ReportRenderError("session canonical hash does not match session content")
    if report.session_manifest_id != session_manifest.session_manifest_id:
        raise ReportRenderError("report and session manifest IDs do not match")
    if report.session_manifest_hash != session_manifest.canonical_hash:
        raise ReportRenderError("report and session manifest hashes do not match")
    if report.governance_snapshot_id != session_manifest.final_snapshot_id:
        raise ReportRenderError("report is not bound to the session final snapshot")
    if report.generation_status != session_manifest.generation_status:
        raise ReportRenderError("report and session generation statuses do not match")
    if session_manifest.report_hash is not None and (
        session_manifest.report_hash != report.canonical_report_hash
    ):
        raise ReportRenderError("report hash does not match the session manifest")
    if tuple(report.citation_ids) != tuple(session_manifest.final_citation_ids):
        raise ReportRenderError("report and session citation indexes do not match")

    sections = tuple(
        GovernanceSectionView(
            section_id=section.section_id,
            title=section.title,
            findings=tuple(
                GovernanceFindingView(
                    finding_id=finding.finding_id,
                    layer=finding.kind,
                    text=finding.text,
                    citation_ids=tuple(finding.citation_ids),
                    uncertainty=finding.uncertainty,
                    source_role=(
                        finding.source_role.value
                        if getattr(finding, "source_role", None) is not None
                        else None
                    ),
                )
                for finding in section.findings
            ),
        )
        for section in report.sections
    )
    technical_checks = tuple(
        GovernanceValidationCheckView(
            check_code=check.check_code,
            passed=check.passed,
            detail=check.detail,
        )
        for check in report.technical_validation.checks
    )
    return GovernanceReportView(
        schema_version=report.schema_version,
        renderer_version=renderer_version,
        report_id=report.governance_report_id,
        report_hash=report.canonical_report_hash,
        company_id=report.company_id,
        state_at=report.state_at,
        known_at=report.known_at,
        perspective=report.perspective.value,
        snapshot_id=report.governance_snapshot_id,
        snapshot_hash=report.governance_snapshot_hash,
        evidence_manifest_id=report.evidence_manifest_id,
        evidence_manifest_hash=report.evidence_manifest_hash,
        session_manifest_id=report.session_manifest_id,
        session_manifest_hash=report.session_manifest_hash,
        generation_status=report.generation_status.value,
        decision_author=report.decision_author,
        future_knowledge_used=report.future_knowledge_used,
        sections=sections,
        citation_ids=tuple(report.citation_ids),
        active_gap_ids=tuple(report.active_gap_ids),
        active_conflict_ids=tuple(report.active_conflict_ids),
        pending_candidate_ids=tuple(report.pending_candidate_ids),
        data_limitations=tuple(report.data_limitations),
        technical_validation_passed=report.technical_validation.passed,
        hard_failure_codes=tuple(report.technical_validation.hard_failure_codes),
        technical_checks=technical_checks,
        failure_code=report.failure_code,
        session_parent_report_run_id=session_manifest.parent_report_run_id,
        session_model_profile=session_manifest.model_profile,
        session_runner_protocol_version=session_manifest.runner_protocol_version,
        session_tool_protocol_version=session_manifest.tool_protocol_version,
        session_input_pack_id=session_manifest.input_pack_id,
        session_input_pack_hash=session_manifest.input_pack_hash,
        session_initial_snapshot_id=session_manifest.initial_snapshot_id,
        session_final_snapshot_id=session_manifest.final_snapshot_id,
        session_tool_read_ids=tuple(session_manifest.tool_read_ids),
        session_research_task_ids=tuple(session_manifest.research_task_ids),
        session_research_result_bundle_ids=tuple(
            session_manifest.research_result_bundle_ids
        ),
        session_snapshot_adoption_ids=tuple(session_manifest.snapshot_adoption_ids),
        session_final_citation_ids=tuple(session_manifest.final_citation_ids),
    )


_LAYER_LABELS = {
    "fact": "正式事实",
    "contextual": "外部背景",
    "judgment": "Codex 判断",
}


def _display(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return canonical_datetime(value)
    if isinstance(value, bool):
        return "true" if value else "false"
    return redact_text(str(value))


def _markdown_inline(value: object | None) -> str:
    text = _display(value).replace("\r", " ").replace("\n", " ")
    for character in ("\\", "`", "*", "_", "{", "}", "[", "]", "<", ">", "#", "|"):
        text = text.replace(character, f"\\{character}")
    return text


def _join_ids(values: tuple[str, ...]) -> str:
    return ", ".join(_display(value) for value in values) if values else "none"


def render_markdown(view: GovernanceReportView) -> bytes:
    lines = [
        "# 治理与管理层报告",
        "",
        f"> renderer: {_markdown_inline(MARKDOWN_RENDERER_VERSION)}",
        "",
        "| 字段 | 值 |",
        "|---|---|",
    ]
    metadata = (
        ("report_id", view.report_id),
        ("report_hash", view.report_hash),
        ("company_id", view.company_id),
        ("state_at", view.state_at),
        ("known_at", view.known_at),
        ("perspective", view.perspective),
        ("snapshot_id", view.snapshot_id),
        ("snapshot_hash", view.snapshot_hash),
        ("evidence_manifest_id", view.evidence_manifest_id),
        ("evidence_manifest_hash", view.evidence_manifest_hash),
        ("session_manifest_id", view.session_manifest_id),
        ("session_manifest_hash", view.session_manifest_hash),
        ("generation_status", view.generation_status),
        ("decision_author", view.decision_author),
        ("future_knowledge_used", view.future_knowledge_used),
    )
    lines.extend(
        f"| {_markdown_inline(key)} | {_markdown_inline(value)} |"
        for key, value in metadata
    )

    for section in view.sections:
        lines.extend(("", f"## {_markdown_inline(section.title)}", ""))
        if not section.findings:
            lines.append("暂无分节判断。")
            continue
        for finding in section.findings:
            layer = _LAYER_LABELS.get(finding.layer, finding.layer)
            lines.append(
                f"- **{_markdown_inline(layer)}** `{_markdown_inline(finding.finding_id)}`："
                f"{_markdown_inline(finding.text)}"
            )
            lines.append(f"  - 引用：{_markdown_inline(_join_ids(finding.citation_ids))}")
            if finding.uncertainty:
                lines.append(f"  - 不确定性：{_markdown_inline(finding.uncertainty)}")
            if finding.source_role:
                lines.append(f"  - 来源角色：{_markdown_inline(finding.source_role)}")

    lines.extend(
        (
            "",
            "## 不确定性与引用索引",
            "",
            f"- 活动缺口：{_markdown_inline(_join_ids(view.active_gap_ids))}",
            f"- 活动冲突：{_markdown_inline(_join_ids(view.active_conflict_ids))}",
            f"- 待定候选：{_markdown_inline(_join_ids(view.pending_candidate_ids))}",
            f"- 数据限制：{_markdown_inline(_join_ids(view.data_limitations))}",
            f"- 引用：{_markdown_inline(_join_ids(view.citation_ids))}",
            "",
            "## 技术校验",
            "",
            f"- passed：{_markdown_inline(view.technical_validation_passed)}",
            f"- hard failures：{_markdown_inline(_join_ids(view.hard_failure_codes))}",
            f"- failure_code：{_markdown_inline(view.failure_code or 'none')}",
        )
    )
    for check in view.technical_checks:
        suffix = f" — {_markdown_inline(check.detail)}" if check.detail else ""
        lines.append(
            f"- `{_markdown_inline(check.check_code)}`："
            f"{'passed' if check.passed else 'failed'}{suffix}"
        )
    lines.extend(
        (
            "",
            "## 会话",
            "",
            f"- input pack：{_markdown_inline(view.session_input_pack_id)} "
            f"({_markdown_inline(view.session_input_pack_hash)})",
            f"- initial snapshot：{_markdown_inline(view.session_initial_snapshot_id)}",
            f"- final snapshot：{_markdown_inline(view.session_final_snapshot_id)}",
            f"- tool reads：{_markdown_inline(_join_ids(view.session_tool_read_ids))}",
            f"- research tasks：{_markdown_inline(_join_ids(view.session_research_task_ids))}",
            f"- snapshot adoptions：{_markdown_inline(_join_ids(view.session_snapshot_adoption_ids))}",
            "",
        )
    )
    return "\n".join(lines).encode("utf-8", errors="strict")


def _html_list(values: tuple[str, ...]) -> str:
    if not values:
        return '<span class="empty">none</span>'
    return "<ul>" + "".join(f"<li>{html.escape(_display(value))}</li>" for value in values) + "</ul>"


def render_html(view: GovernanceReportView) -> bytes:
    metadata = (
        ("report_id", view.report_id),
        ("report_hash", view.report_hash),
        ("company_id", view.company_id),
        ("state_at", view.state_at),
        ("known_at", view.known_at),
        ("perspective", view.perspective),
        ("snapshot_id", view.snapshot_id),
        ("snapshot_hash", view.snapshot_hash),
        ("evidence_manifest_id", view.evidence_manifest_id),
        ("evidence_manifest_hash", view.evidence_manifest_hash),
        ("session_manifest_id", view.session_manifest_id),
        ("session_manifest_hash", view.session_manifest_hash),
        ("generation_status", view.generation_status),
        ("decision_author", view.decision_author),
        ("future_knowledge_used", view.future_knowledge_used),
    )
    metadata_html = "".join(
        f"<tr><th>{html.escape(key)}</th><td>{html.escape(_display(value))}</td></tr>"
        for key, value in metadata
    )
    sections_html: list[str] = []
    for section in view.sections:
        findings_html = "".join(
            (
                f'<li class="finding {html.escape(finding.layer)}">'
                f'<span class="layer">{html.escape(_LAYER_LABELS.get(finding.layer, finding.layer))}</span>'
                f'<span class="finding-id">{html.escape(_display(finding.finding_id))}</span>'
                f'<p>{html.escape(_display(finding.text))}</p>'
                f'<div class="citations">引用：{html.escape(_join_ids(finding.citation_ids))}</div>'
                + (
                    f'<div class="uncertainty">不确定性：{html.escape(_display(finding.uncertainty))}</div>'
                    if finding.uncertainty
                    else ""
                )
                + (
                    f'<div class="source-role">来源角色：{html.escape(_display(finding.source_role))}</div>'
                    if finding.source_role
                    else ""
                )
                + "</li>"
            )
            for finding in section.findings
        )
        if not findings_html:
            findings_html = '<li class="empty">暂无分节判断。</li>'
        sections_html.append(
            f'<section data-section-id="{html.escape(_display(section.section_id), quote=True)}">'
            f"<h2>{html.escape(_display(section.title))}</h2>"
            f'<ul class="findings">{findings_html}</ul></section>'
        )
    checks_html = "".join(
        f'<li data-passed="{str(check.passed).lower()}"><code>{html.escape(_display(check.check_code))}</code>'
        f": {'passed' if check.passed else 'failed'}"
        + (f" — {html.escape(_display(check.detail))}" if check.detail else "")
        + "</li>"
        for check in view.technical_checks
    )
    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>治理与管理层报告 {html.escape(_display(view.report_id))}</title>
<style>
body{{font-family:system-ui,-apple-system,"Segoe UI",sans-serif;line-height:1.55;color:#17202a;max-width:1100px;margin:2rem auto;padding:0 1rem}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ccd1d1;padding:.45rem;text-align:left;vertical-align:top}}th{{width:16rem;background:#f4f6f7}}
.layer{{font-weight:700;margin-right:.7rem}}.finding-id{{font-family:monospace;color:#566573}}.empty{{color:#6c757d}}code{{font-family:monospace}}
</style>
</head>
<body data-renderer-version="{html.escape(HTML_RENDERER_VERSION, quote=True)}" data-report-hash="{html.escape(_display(view.report_hash), quote=True)}" data-snapshot-id="{html.escape(_display(view.snapshot_id), quote=True)}">
<h1>治理与管理层报告</h1>
<table><tbody>{metadata_html}</tbody></table>
{''.join(sections_html)}
<section><h2>不确定性与引用索引</h2>
<h3>活动缺口</h3>{_html_list(view.active_gap_ids)}
<h3>活动冲突</h3>{_html_list(view.active_conflict_ids)}
<h3>待定候选</h3>{_html_list(view.pending_candidate_ids)}
<h3>数据限制</h3>{_html_list(view.data_limitations)}
<h3>引用</h3>{_html_list(view.citation_ids)}
</section>
<section><h2>技术校验</h2><p>passed: {str(view.technical_validation_passed).lower()}</p><p>hard failures: {html.escape(_join_ids(view.hard_failure_codes))}</p><p>failure_code: {html.escape(_display(view.failure_code or 'none'))}</p><ul>{checks_html}</ul></section>
<section><h2>会话</h2><p>input pack: {html.escape(_display(view.session_input_pack_id))} ({html.escape(_display(view.session_input_pack_hash))})</p><p>initial snapshot: {html.escape(_display(view.session_initial_snapshot_id))}</p><p>final snapshot: {html.escape(_display(view.session_final_snapshot_id))}</p></section>
</body>
</html>
"""
    return document.encode("utf-8", errors="strict")


def _excel_value(value: object | None) -> str:
    rendered = _display(value)
    if rendered.startswith(("=", "+", "-", "@")):
        return "'" + rendered
    return rendered


def _write_sheet(workbook: Workbook, title: str, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        worksheet.append(tuple(_excel_value(value) for value in row))
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(max_length + 2, 12), 70
        )


def _deterministic_workbook_bytes(workbook: Workbook) -> bytes:
    # openpyxl.save() replaces modified with wall-clock time. Preserve the
    # explicitly chosen properties before saving, then freeze those bytes too.
    core_properties = tostring(workbook.properties.to_tree())
    raw = BytesIO()
    workbook.save(raw)
    normalized = BytesIO()
    with ZipFile(BytesIO(raw.getvalue()), "r") as source, ZipFile(
        normalized, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(filename=name, date_time=_FIXED_ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            payload = core_properties if name == "docProps/core.xml" else source.read(name)
            target.writestr(info, payload, compress_type=ZIP_DEFLATED, compresslevel=9)
    return normalized.getvalue()


def render_xlsx(view: GovernanceReportView) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.iso_dates = True
    workbook.properties.creator = "AstraValue"
    workbook.properties.lastModifiedBy = "AstraValue"
    workbook.properties.created = _FIXED_WORKBOOK_TIMESTAMP
    workbook.properties.modified = _FIXED_WORKBOOK_TIMESTAMP
    workbook.properties.title = "治理与管理层报告"
    workbook.properties.subject = _display(view.report_id)
    workbook.properties.keywords = _display(view.report_hash)

    summary_rows = [
        ("renderer_version", XLSX_RENDERER_VERSION),
        ("report_id", view.report_id),
        ("report_hash", view.report_hash),
        ("company_id", view.company_id),
        ("state_at", view.state_at),
        ("known_at", view.known_at),
        ("perspective", view.perspective),
        ("snapshot_id", view.snapshot_id),
        ("snapshot_hash", view.snapshot_hash),
        ("evidence_manifest_id", view.evidence_manifest_id),
        ("evidence_manifest_hash", view.evidence_manifest_hash),
        ("session_manifest_id", view.session_manifest_id),
        ("session_manifest_hash", view.session_manifest_hash),
        ("generation_status", view.generation_status),
        ("decision_author", view.decision_author),
        ("future_knowledge_used", view.future_knowledge_used),
        ("pending_candidate_ids", _join_ids(view.pending_candidate_ids)),
        ("data_limitations", _join_ids(view.data_limitations)),
        ("technical_validation_passed", view.technical_validation_passed),
        ("hard_failure_codes", _join_ids(view.hard_failure_codes)),
        ("failure_code", view.failure_code or "none"),
    ]
    _write_sheet(workbook, "Summary", ("Field", "Value"), summary_rows)

    finding_rows: list[tuple[object, ...]] = []
    for section in view.sections:
        for finding in section.findings:
            finding_rows.append(
                (
                    section.section_id,
                    section.title,
                    finding.finding_id,
                    finding.layer,
                    finding.text,
                    _join_ids(finding.citation_ids),
                    finding.uncertainty or "",
                    finding.source_role or "",
                )
            )
    _write_sheet(
        workbook,
        "Findings",
        (
            "Section ID",
            "Section",
            "Finding ID",
            "Layer",
            "Text",
            "Citation IDs",
            "Uncertainty",
            "Source Role",
        ),
        finding_rows,
    )
    _write_sheet(
        workbook,
        "Evidence",
        ("Citation ID", "Report ID", "Snapshot ID", "Manifest ID"),
        [
            (citation_id, view.report_id, view.snapshot_id, view.evidence_manifest_id)
            for citation_id in view.citation_ids
        ],
    )
    _write_sheet(
        workbook,
        "Gaps",
        ("Gap ID", "Snapshot ID"),
        [(gap_id, view.snapshot_id) for gap_id in view.active_gap_ids],
    )
    _write_sheet(
        workbook,
        "Conflicts",
        ("Conflict ID", "Snapshot ID"),
        [
            (conflict_id, view.snapshot_id)
            for conflict_id in view.active_conflict_ids
        ],
    )
    session_rows = [
        ("session_manifest_id", view.session_manifest_id),
        ("session_manifest_hash", view.session_manifest_hash),
        ("parent_report_run_id", view.session_parent_report_run_id),
        ("model_profile", view.session_model_profile),
        ("runner_protocol_version", view.session_runner_protocol_version),
        ("tool_protocol_version", view.session_tool_protocol_version),
        ("input_pack_id", view.session_input_pack_id),
        ("input_pack_hash", view.session_input_pack_hash),
        ("initial_snapshot_id", view.session_initial_snapshot_id),
        ("final_snapshot_id", view.session_final_snapshot_id),
        ("tool_read_ids", _join_ids(view.session_tool_read_ids)),
        ("research_task_ids", _join_ids(view.session_research_task_ids)),
        (
            "research_result_bundle_ids",
            _join_ids(view.session_research_result_bundle_ids),
        ),
        ("snapshot_adoption_ids", _join_ids(view.session_snapshot_adoption_ids)),
        ("final_citation_ids", _join_ids(view.session_final_citation_ids)),
    ]
    _write_sheet(workbook, "Session", ("Field", "Value"), session_rows)
    return _deterministic_workbook_bytes(workbook)


HtmlToPdf = Callable[[bytes], bytes]


def render_pdf(
    view: GovernanceReportView,
    *,
    html_to_pdf: HtmlToPdf,
) -> bytes:
    """Render PDF through an injected converter using exactly the canonical HTML view."""

    if not callable(html_to_pdf):
        raise ReportRenderError("html_to_pdf must be callable")
    html_bytes = render_html(view)
    rendered = html_to_pdf(html_bytes)
    if not isinstance(rendered, (bytes, bytearray, memoryview)):
        raise ReportRenderError("html_to_pdf must return bytes")
    pdf_bytes = bytes(rendered)
    if not pdf_bytes.startswith(b"%PDF-"):
        raise ReportRenderError("html_to_pdf returned an invalid PDF payload")
    return pdf_bytes


__all__ = [
    "GovernanceFindingView",
    "GovernanceReportView",
    "GovernanceSectionView",
    "GovernanceValidationCheckView",
    "HtmlToPdf",
    "ReportRenderError",
    "build_governance_report_view",
    "render_html",
    "render_markdown",
    "render_pdf",
    "render_xlsx",
]
