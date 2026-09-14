from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from analysis.exports import export_report, render_html, render_markdown
from analysis.models import (
    FactRecord,
    ReportCreateRequest,
    SourceRecord,
    StructuredAcquisitionMethod,
    StructuredFactAdmission,
    StructuredFactNature,
    StructuredQualityStatus,
    VerificationStatus,
)
from analysis.reporting import ReportBuilder
from analysis.structured.reporting_bridge import build_report_request
from analysis.structured.scope import LITE_PROFILE_ID, load_research_profile
from analysis.structured.storage import canonical_json


NOW = datetime(2026, 9, 13, 4, tzinfo=timezone.utc)


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(canonical_json(value) + "\n", encoding="utf-8")


def _fact(metric: str, period: date, period_type: str, value: float, index: int) -> FactRecord:
    snapshot_hash = hashlib.sha256(f"snapshot-{index}".encode()).hexdigest()
    source_id = f"source-{index}"
    admission = StructuredFactAdmission(
        acquisition_method=StructuredAcquisitionMethod.SUPPLIER_STRUCTURED,
        nature=StructuredFactNature.OBSERVED,
        quality=StructuredQualityStatus.PASSED,
        source_policy_id="structured-test",
        source_policy_version="1.0.0",
        field_definition_id=f"test.{metric}",
        field_definition_version="1.0.0",
        raw_resource_snapshot_id=f"snapshot-{index}",
        snapshot_sha256=snapshot_hash,
        row_key=f"row-{index}",
        field_path=f"$.{metric}",
        original_value=value,
        original_unit="CNY",
        retrieved_at=NOW,
        available_at=NOW,
    )
    return FactRecord(
        fact_id=f"fact-{index}",
        ticker="600519",
        metric_id=metric,
        value=value,
        unit="CNY",
        period_start=date(period.year, 1, 1) if period_type == "cumulative" else None,
        period_end=period,
        period_type=period_type,
        as_of=NOW,
        source_ids=[source_id],
        verification_status=VerificationStatus.SUPPLIER_DIRECT,
        metadata={
            "requires_materialization_selection": True,
            "available_at": NOW.isoformat(),
            "decimal_value": str(value),
            "source_definition_id": "structured-test",
            "source_definition_version": "1.0.0",
        },
        structured_admission=admission,
    )


def _pack(tmp_path: Path) -> Path:
    annual = [date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)]
    quarters = [
        date(2024, 9, 30),
        date(2024, 12, 31),
        date(2025, 3, 31),
        date(2025, 6, 30),
        date(2025, 9, 30),
        date(2025, 12, 31),
        date(2026, 3, 31),
        date(2026, 6, 30),
    ]
    facts = [
        *[
            _fact("operating_income", period, "cumulative", 1000 + index, index)
            for index, period in enumerate(annual, start=1)
        ],
        *[
            _fact("operating_income", period, "single_quarter", 100 + index, index)
            for index, period in enumerate(quarters, start=10)
        ],
        _fact("operating_cost", annual[-1], "cumulative", 300, 30),
        _fact("market_price", date(2026, 9, 13), "current", 1234.56, 31),
    ]
    projection = tmp_path / "projection" / "600519"
    projection.mkdir(parents=True)
    facts_path = projection / "coverage-facts.jsonl"
    facts_path.write_text(
        "".join(canonical_json(item.model_dump(mode="json")) + "\n" for item in facts),
        encoding="utf-8",
    )
    sources = []
    for fact in facts:
        admission = fact.structured_admission
        assert admission is not None
        sources.append(
            SourceRecord(
                source_id=fact.source_ids[0],
                name="测试结构化响应",
                source_type="supplier-structured",
                upstream_source_id="test",
                retrieved_at=NOW,
                document_hash=admission.snapshot_sha256,
                source_definition_id="structured-test",
                source_definition_version="1.0.0",
                raw_resource_snapshot_id=admission.raw_resource_snapshot_id,
                available_at=NOW,
            )
        )
    (projection / "sources.jsonl").write_text(
        "".join(canonical_json(item.model_dump(mode="json")) + "\n" for item in sources),
        encoding="utf-8",
    )
    projection_manifest = projection / "manifest.json"
    _write_json(
        projection_manifest,
        {
            "run_id": "run-test",
            "source_namespace_id": "namespace-test",
            "contract_hash": "a" * 64,
            "materialization_hash": "b" * 64,
        },
    )

    profile = load_research_profile(LITE_PROFILE_ID)
    periods = {
        "annual": [item.isoformat() for item in annual],
        "quarters": [item.isoformat() for item in quarters],
        "dependencies": [],
        "current": "2026-09-13",
    }
    pack = tmp_path / "pack"
    pack.mkdir()
    core = {
        "pack_id": "lite-pack-test",
        "profile_id": profile["profile_id"],
        "profile_sha256": profile["content_sha256"],
        "ticker": "600519",
        "company_name": "贵州茅台",
        "as_of": "2026-09-13",
        "status": "ready",
        "periods": periods,
        "metrics": [
            {
                "state": "ready",
                "fact": {
                    **item.model_dump(mode="json"),
                    "available_at": NOW.isoformat(),
                },
            }
            for item in facts
        ],
    }
    coverage = {
        "profile_id": profile["profile_id"],
        "profile_sha256": profile["content_sha256"],
        "ticker": "600519",
        "as_of": "2026-09-13",
        "periods": periods,
        "requirements": [
            {
                "requirement_id": f"requirement-{item.fact_id}",
                "state": "ready",
                "fact_ids": [item.fact_id],
            }
            for item in facts
        ],
        "question_routes": [
            {
                "question_id": "ES02.Q01",
                "title": "历史财务表现如何？",
                "quality_state": "ready",
                "scope_route": "core",
                "activation_status": "active",
                "full_requirement_refs": ["REQ.ES02.Q01.001"],
                "quality_counts": {"ready": len(facts)},
            }
        ],
        "counts": {"ready": len(facts)},
        "question_quality_counts": {"ready": 1},
        "full_coverage_preserved": True,
    }
    _write_json(pack / "core-pack.json", core)
    _write_json(pack / "core-coverage.json", coverage)
    manifest = {
        "pack_id": "lite-pack-test",
        "pack_identity_hash": "c" * 64,
        "pack_version": "eight-step-lite-pack-test",
        "profile_id": profile["profile_id"],
        "profile_sha256": profile["content_sha256"],
        "ticker": "600519",
        "as_of": "2026-09-13",
        "status": "ready",
        "periods": periods,
        "source_inputs": [
            {
                "company_path": str(projection.resolve()),
                "files": {
                    "coverage-facts.jsonl": {
                        "path": str(facts_path.resolve()),
                        "sha256": _hash(facts_path),
                    }
                },
                "manifests": [
                    {
                        "path": str(projection_manifest.resolve()),
                        "sha256": _hash(projection_manifest),
                    }
                ],
            }
        ],
        "auxiliary_inputs": [],
        "peer_inputs": [],
        "output_hashes": {
            "core-pack.json": _hash(pack / "core-pack.json"),
            "core-coverage.json": _hash(pack / "core-coverage.json"),
        },
    }
    _write_json(pack / "manifest.json", manifest)
    return pack


def test_lite_pack_builds_report_request_with_selection_aliases_and_period_targets(tmp_path):
    request = build_report_request(_pack(tmp_path))
    assert request.ticker == "600519"
    assert request.industry == "消费"
    assert request.current_price == 1234.56
    assert request.materialization_selected_fact_ids
    assert all(item.structured_admission for item in request.facts if item.fact_id.startswith("fact-"))

    report = ReportBuilder().build(request)
    assert report.request_metadata["industry_route"] == "consumer"
    assert report.conclusion.rating.value == "暂不评级"
    assert {item.status.value for item in report.audit.model_runs} == {"待核验"}
    assert [(item["check"], item["status"]) for item in report.audit.data_quality_checks[:2]] == [
        ("3年年度覆盖", "OK"),
        ("8个单季覆盖", "OK"),
    ]
    revenue = [item for item in report.facts if item.metric_id == "revenue"]
    assert len([item for item in revenue if item.period_type == "annual"]) == 3
    assert len([item for item in revenue if item.period_type == "single_quarter"]) == 8
    assert all(item.derived_from_fact_ids for item in revenue)


def test_pack_hash_change_fails_closed(tmp_path):
    pack = _pack(tmp_path)
    with (pack / "core-coverage.json").open("a", encoding="utf-8") as stream:
        stream.write(" ")
    with pytest.raises(ValueError, match="report_pack_output_hash_mismatch"):
        build_report_request(pack)


def test_budget_exceeded_pack_is_not_reportable(tmp_path):
    pack = _pack(tmp_path)
    core_path = pack / "core-pack.json"
    core = json.loads(core_path.read_text(encoding="utf-8"))
    core["status"] = "budget_exceeded"
    _write_json(core_path, core)
    manifest_path = pack / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["status"] = "budget_exceeded"
    manifest["output_hashes"]["core-pack.json"] = _hash(core_path)
    _write_json(manifest_path, manifest)

    with pytest.raises(ValueError, match="report_pack_status_not_reportable:budget_exceeded"):
        build_report_request(pack)


def test_report_system_metadata_cannot_be_overridden():
    report = ReportBuilder().build(
        ReportCreateRequest(
            ticker="600519",
            company_name="贵州茅台",
            industry="消费",
            as_of=NOW,
            report_notes="系统报告说明",
            sync_result_id="system-sync",
            input_metadata={
                "industry_route": "bank",
                "report_notes": "伪造报告说明",
                "sync_result_id": "forged-sync",
                "custom_key": "retained",
            },
        )
    )

    assert report.request_metadata == {
        "custom_key": "retained",
        "industry_route": "consumer",
        "report_notes": "系统报告说明",
        "sync_result_id": "system-sync",
        "dimensional_sync_result_id": None,
        "event_sync_result_id": None,
    }


def test_bank_and_unknown_industry_never_use_general_enterprise_models():
    bank = ReportBuilder().build(
        ReportCreateRequest(
            ticker="600000",
            company_name="银行样例",
            industry="银行",
            as_of=NOW,
        )
    )
    bank_methods = {item.method_ref.split("@", 1)[0] for item in bank.audit.model_runs}
    assert bank.request_metadata["industry_route"] == "bank"
    assert bank_methods == {"VAL.RESIDUAL_INCOME", "VAL.DDM", "VAL.RELATIVE"}
    assert "VAL.DCF.FCFF" not in bank_methods

    unknown = ReportBuilder().build(
        ReportCreateRequest(
            ticker="000001",
            company_name="未知行业样例",
            industry="未知细分",
            as_of=NOW,
        )
    )
    assert unknown.request_metadata["industry_route"] == "unknown"
    assert unknown.audit.model_runs == []
    assert any("不自动套用普通企业" in item for item in unknown.sections[6].warnings)


def test_report_exports_share_price_snapshot_and_fact_lineage(tmp_path, monkeypatch):
    report = ReportBuilder().build(build_report_request(_pack(tmp_path)))
    markdown = render_markdown(report)
    html = render_html(report)
    assert "1234.56" in markdown
    assert "1234.56" in html
    assert report.data_snapshot_id in markdown and report.data_snapshot_id in html
    assert "fact-1" in markdown and "fact-1" in html

    xlsx = export_report(report, "xlsx", tmp_path / "exports")
    workbook = load_workbook(xlsx, data_only=False)
    assert workbook["置顶结论"]["A10"].value == 1234.56
    assert report.data_snapshot_id in [cell.value for cell in workbook["置顶结论"][7]]
    facts_sheet = workbook["财务事实"]
    assert "fact-1" in [facts_sheet.cell(row, 1).value for row in range(2, facts_sheet.max_row + 1)]

    captured = {}

    def fake_pdf(html_path, pdf_path):
        captured["html"] = html_path.read_text(encoding="utf-8")
        pdf_path.write_bytes(b"%PDF-1.4\nunit-test")

    monkeypatch.setattr("analysis.exports._html_to_pdf", fake_pdf)
    pdf = export_report(report, "pdf", tmp_path / "exports")
    assert pdf.read_bytes().startswith(b"%PDF")
    assert "1234.56" in captured["html"]
    assert report.data_snapshot_id in captured["html"]
