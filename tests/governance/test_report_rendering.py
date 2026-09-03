from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from analysis.governance.codex_tools import canonical_session_manifest_hash
from analysis.governance.models import (
    CodexInputPack,
    CodexJudgment,
    CodexSessionManifest,
    CompletenessStatus,
    ContextualFinding,
    DeltaDisposition,
    FactualFinding,
    GovernancePerspective,
    GovernanceQuestionCoverageLink,
    GovernanceReport,
    GovernanceReportSection,
    GovernanceSnapshot,
    QuestionSummary,
    ReportGenerationStatus,
    ReportTechnicalValidation,
    ResearchBudget,
    SnapshotAnchorLink,
    SnapshotDeltaLink,
    SnapshotRecordLink,
    TimePrecision,
    ToolSchemaReference,
    ValidationCheck,
)
from analysis.governance.renderers import (
    GovernanceReportView,
    ReportRenderError,
    build_governance_report_view,
    render_html,
    render_markdown,
    render_pdf,
    render_xlsx,
)
from analysis.governance.report_service import canonical_governance_report_hash


NOW = datetime(2024, 6, 30, 15, 0, tzinfo=timezone.utc)
H1 = "1" * 64
H2 = "2" * 64
H3 = "3" * 64
H4 = "4" * 64
H5 = "5" * 64
H6 = "6" * 64
QUESTION = "GOV.Q04.EXECUTIVE_ROSTER_BACKGROUND"


@dataclass(frozen=True)
class _ReportBundle:
    report: GovernanceReport
    session: CodexSessionManifest
    snapshot: GovernanceSnapshot
    input_pack: CodexInputPack


def _make_bundle(*, failed: bool = False, sensitive: bool = False) -> _ReportBundle:
    snapshot_id = "govsnapshot:render"
    anchor = SnapshotAnchorLink(
        anchor_link_id="govanchor:roster",
        governance_snapshot_id=snapshot_id,
        question_id=QUESTION,
        state_kind="roster",
        anchor_record_id="govrec:roster",
        reference_at=NOW,
        available_at=NOW,
        canonical_hash=H1,
    )
    applied = SnapshotDeltaLink(
        delta_link_id="govdelta:appointment",
        governance_snapshot_id=snapshot_id,
        question_id=QUESTION,
        delta_record_id="govrec:appointment",
        disposition=DeltaDisposition.APPLIED,
        sequence=1,
        effective_at=NOW,
        available_at=NOW,
        canonical_hash=H2,
    )
    excluded = SnapshotDeltaLink(
        delta_link_id="govdelta:future",
        governance_snapshot_id=snapshot_id,
        question_id=QUESTION,
        delta_record_id="govrec:future",
        disposition=DeltaDisposition.EXCLUDED,
        sequence=2,
        effective_at=NOW,
        available_at=NOW,
        exclusion_reason="superseded_by_visible_correction",
        canonical_hash=H3,
    )
    record_link = SnapshotRecordLink(
        record_link_id="govrecordlink:tenure",
        governance_snapshot_id=snapshot_id,
        record_id="govrec:tenure",
        record_kind="role_tenure",
        canonical_hash=H1,
    )
    snapshot = GovernanceSnapshot(
        governance_snapshot_id=snapshot_id,
        company_id="cn-600519",
        state_at=NOW,
        known_at=NOW,
        state_time_precision=TimePrecision.DATETIME,
        known_time_precision=TimePrecision.DATETIME,
        perspective=GovernancePerspective.STRICT,
        question_set_version="1.0.0",
        source_registry_version="sources:1",
        query_pack_version="queries:1",
        extractor_versions=("tables:1",),
        reconstruction_version="reconstruction:1",
        evidence_manifest_id="shared-manifest:render",
        evidence_manifest_hash=H2,
        anchor_links=(anchor,),
        delta_links=(applied, excluded),
        record_links=(record_link,),
        canonical_record_ids=("govrec:tenure",),
        active_gap_ids=("govgap:missing-roster-history",),
        active_conflict_ids=("govconflict:overlap",),
        pending_candidate_ids=("govclaim:pending",),
        question_level_coverage=(
            GovernanceQuestionCoverageLink(
                question_id=QUESTION,
                coverage_entry_ids=("shared-coverage:one",),
                completeness_status=CompletenessStatus.CONFLICTED,
            ),
        ),
        completeness_status=CompletenessStatus.CONFLICTED,
        future_knowledge_used=False,
        canonical_snapshot_hash=H3,
        created_at=NOW,
    )
    input_pack = CodexInputPack(
        input_pack_id="govinput:render",
        company_id=snapshot.company_id,
        state_at=NOW,
        known_at=NOW,
        perspective=GovernancePerspective.STRICT,
        governance_snapshot_id=snapshot_id,
        governance_snapshot_hash=H3,
        evidence_manifest_id=snapshot.evidence_manifest_id,
        evidence_manifest_hash=snapshot.evidence_manifest_hash,
        question_set_version="1.0.0",
        question_summaries=(
            QuestionSummary(
                question_id=QUESTION,
                completeness_status=CompletenessStatus.CONFLICTED,
                coverage_entry_ids=("shared-coverage:one",),
                anchor_record_ids=("govrec:roster",),
                summary="名册锚点存在，但任期仍有冲突。",
            ),
        ),
        active_gap_ids=snapshot.active_gap_ids,
        active_conflict_ids=snapshot.active_conflict_ids,
        pending_candidate_ids=snapshot.pending_candidate_ids,
        tool_schemas=(
            ToolSchemaReference(
                tool_name="get_governance_overview",
                tool_version="1.0.0",
                input_schema_hash=H1,
                output_schema_hash=H2,
            ),
        ),
        research_budget=ResearchBudget(
            max_rounds=2,
            max_child_tasks=2,
            max_network_requests=4,
            max_parallelism=1,
            wall_clock_seconds=30,
            max_output_bytes=100_000,
        ),
        temporal_rules=("available_at_lte_known_at",),
        report_output_schema_hash=H4,
        canonical_hash=H4,
        created_at=NOW,
    )
    report_hash = H5
    session_hash = H6
    generation_status = (
        ReportGenerationStatus.FAILED if failed else ReportGenerationStatus.COMPLETED
    )
    citations = ("citation:context", "citation:fact", "citation:judgment")
    session = CodexSessionManifest(
        session_manifest_id="govsession:render",
        parent_report_run_id="report-run:render",
        model_profile="codex-governance-v1",
        runner_protocol_version="1.0.0",
        tool_protocol_version="1.0.0",
        input_pack_id=input_pack.input_pack_id,
        input_pack_hash=input_pack.canonical_hash,
        initial_snapshot_id=snapshot_id,
        final_snapshot_id=snapshot_id,
        tool_read_ids=("govtoolread:one",),
        research_task_ids=("govresearchtask:one",),
        research_result_bundle_ids=("govresearchbundle:one",),
        final_citation_ids=citations,
        report_hash=report_hash,
        generation_status=generation_status,
        output_schema_validated=not failed,
        failure_code="future_leakage" if failed else None,
        started_at=NOW,
        completed_at=NOW,
        canonical_hash=session_hash,
    )
    session_hash = canonical_session_manifest_hash(session)
    session = session.model_copy(update={"canonical_hash": session_hash})
    sensitive_text = (
        "Authorization: Bearer should-not-appear\n"
        "证据缓存位于 D:\\Private Folder\\evidence.json，"
        "详情 https://example.invalid/path?token=should-not-appear"
    )
    checks = (
        ValidationCheck(
            check_code="future_firewall",
            passed=not failed,
            detail=sensitive_text if sensitive else "strict 时点门已执行",
        ),
    )
    technical_validation = ReportTechnicalValidation(
        passed=not failed,
        hard_failure_codes=("future_leakage",) if failed else (),
        checks=checks,
    )
    report = GovernanceReport(
        governance_report_id="govreport:render",
        company_id=snapshot.company_id,
        state_at=NOW,
        known_at=NOW,
        perspective=GovernancePerspective.STRICT,
        governance_snapshot_id=snapshot_id,
        governance_snapshot_hash=H3,
        evidence_manifest_id=snapshot.evidence_manifest_id,
        evidence_manifest_hash=snapshot.evidence_manifest_hash,
        session_manifest_id=session.session_manifest_id,
        session_manifest_hash=session_hash,
        generation_status=generation_status,
        sections=(
            GovernanceReportSection(
                section_id="people",
                title="董事会与管理层",
                findings=(
                    FactualFinding(
                        finding_id="govfinding:fact",
                        text="正式名册披露一名董事。",
                        citation_ids=("citation:fact",),
                    ),
                    ContextualFinding(
                        finding_id="govfinding:context",
                        text="公开背景材料提供补充语境。",
                        citation_ids=("citation:context",),
                        uncertainty="非正式材料不进入权威画像",
                    ),
                    CodexJudgment(
                        finding_id="govfinding:judgment",
                        text="Codex 认为任期稳定性仍需持续观察。",
                        citation_ids=("citation:judgment",),
                        uncertainty="历史材料存在缺口",
                    ),
                ),
            ),
        ),
        active_gap_ids=snapshot.active_gap_ids,
        active_conflict_ids=snapshot.active_conflict_ids,
        pending_candidate_ids=snapshot.pending_candidate_ids,
        data_limitations=(sensitive_text,) if sensitive else ("limited_history",),
        citation_ids=citations,
        future_knowledge_used=False,
        technical_validation=technical_validation,
        failure_code="future_leakage" if failed else None,
        created_at=NOW,
        canonical_report_hash=report_hash,
    )
    report_hash = canonical_governance_report_hash(report)
    report = report.model_copy(update={"canonical_report_hash": report_hash})
    session = session.model_copy(update={"report_hash": report_hash})
    return _ReportBundle(report, session, snapshot, input_pack)


@pytest.fixture
def canonical_view() -> GovernanceReportView:
    bundle = _make_bundle()
    return build_governance_report_view(bundle.report, bundle.session)


def test_canonical_view_fixes_identifiers_hashes_and_sections(
    canonical_view: GovernanceReportView,
) -> None:
    assert canonical_view.report_id == "govreport:render"
    assert canonical_view.report_hash != canonical_view.session_manifest_hash
    assert canonical_view.session_manifest_id == "govsession:render"
    assert canonical_view.snapshot_id == "govsnapshot:render"
    assert canonical_view.snapshot_hash == H3
    assert canonical_view.evidence_manifest_id == "shared-manifest:render"
    assert canonical_view.state_at == canonical_view.known_at == NOW
    assert [section.section_id for section in canonical_view.sections] == ["people"]
    assert {finding.layer for finding in canonical_view.sections[0].findings} == {
        "fact",
        "contextual",
        "judgment",
    }
    assert canonical_view.canonical_view_hash == canonical_view.canonical_view_hash


def test_canonical_view_rejects_cross_session_identity() -> None:
    bundle = _make_bundle()
    altered = bundle.session.model_copy(
        update={"session_manifest_id": "govsession:different"}
    )
    altered = altered.model_copy(
        update={"canonical_hash": canonical_session_manifest_hash(altered)}
    )
    with pytest.raises(ReportRenderError, match="IDs do not match"):
        build_governance_report_view(bundle.report, altered)


def test_canonical_view_recomputes_report_and_session_hashes() -> None:
    bundle = _make_bundle()
    changed_report = bundle.report.model_copy(update={"company_id": "cn-000001"})
    with pytest.raises(ReportRenderError, match="report canonical hash"):
        build_governance_report_view(changed_report, bundle.session)

    changed_session = bundle.session.model_copy(update={"model_profile": "tampered"})
    with pytest.raises(ReportRenderError, match="session canonical hash"):
        build_governance_report_view(bundle.report, changed_session)


def test_markdown_and_html_render_the_same_citations_without_reanalysis(
    canonical_view: GovernanceReportView,
) -> None:
    markdown = render_markdown(canonical_view).decode("utf-8")
    html = render_html(canonical_view).decode("utf-8")
    for citation_id in canonical_view.citation_ids:
        assert citation_id in markdown
        assert citation_id in html
    assert "正式事实" in markdown and "正式事实" in html
    assert "外部背景" in markdown and "外部背景" in html
    assert "Codex 判断" in markdown and "Codex 判断" in html
    assert canonical_view.report_hash in markdown and canonical_view.report_hash in html
    assert canonical_view.snapshot_hash in markdown and canonical_view.snapshot_hash in html


def test_xlsx_has_fixed_sheets_and_the_same_snapshot(
    canonical_view: GovernanceReportView,
) -> None:
    workbook = load_workbook(BytesIO(render_xlsx(canonical_view)), read_only=True)
    assert workbook.sheetnames == [
        "Summary",
        "Findings",
        "Evidence",
        "Gaps",
        "Conflicts",
        "Session",
    ]
    summary = {
        row[0]: row[1]
        for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)
    }
    assert summary["report_hash"] == canonical_view.report_hash
    assert summary["snapshot_id"] == canonical_view.snapshot_id
    assert summary["snapshot_hash"] == canonical_view.snapshot_hash
    assert summary["evidence_manifest_id"] == canonical_view.evidence_manifest_id
    assert summary["session_manifest_id"] == canonical_view.session_manifest_id
    evidence = list(workbook["Evidence"].iter_rows(min_row=2, values_only=True))
    assert {row[0] for row in evidence} == set(canonical_view.citation_ids)


def test_pdf_uses_exactly_the_same_html_view_and_report_hash(
    canonical_view: GovernanceReportView,
) -> None:
    received: list[bytes] = []

    def converter(html_bytes: bytes) -> bytes:
        received.append(html_bytes)
        return b"%PDF-1.7\n" + canonical_view.report_hash.encode("ascii")

    pdf = render_pdf(canonical_view, html_to_pdf=converter)
    assert received == [render_html(canonical_view)]
    assert canonical_view.report_hash.encode("ascii") in pdf
    with pytest.raises(ReportRenderError, match="invalid PDF"):
        render_pdf(canonical_view, html_to_pdf=lambda _html: b"not-a-pdf")


def test_report_outputs_are_utf8_redacted_and_stable() -> None:
    bundle = _make_bundle(sensitive=True)
    view = build_governance_report_view(bundle.report, bundle.session)
    markdown_first = render_markdown(view)
    html_first = render_html(view)
    xlsx_first = render_xlsx(view)
    assert markdown_first == render_markdown(view)
    assert html_first == render_html(view)
    assert xlsx_first == render_xlsx(view)
    markdown_first.decode("utf-8", errors="strict")
    html_first.decode("utf-8", errors="strict")
    combined = markdown_first + html_first
    assert b"should-not-appear" not in combined
    assert b"Private Folder" not in combined
    assert b"[REDACTED]" in combined or b"%5BREDACTED%5D" in combined


def test_pdf_render_is_stable_with_a_deterministic_converter(
    canonical_view: GovernanceReportView,
) -> None:
    converter = lambda html_bytes: b"%PDF-1.7\n" + html_bytes
    assert render_pdf(canonical_view, html_to_pdf=converter) == render_pdf(
        canonical_view, html_to_pdf=converter
    )
