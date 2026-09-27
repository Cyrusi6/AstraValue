"""Real-cache integration; test prose never constitutes host research or golden acceptance."""
import json
import os
from pathlib import Path

import pytest

from analysis.research.calculations import Calculations
from analysis.research.drafts import Drafts
from analysis.research.reports import Reports
from analysis.research.workspace import ResearchError, ResearchWorkspace, sha


@pytest.mark.skipif(os.environ.get("ASTRAVALUE_REAL_CACHE_TEST") != "1", reason="explicit local real-cache opt-in")
def test_real_frozen_pack_to_existing_report_and_exports(tmp_path):
    root = Path(__file__).resolve().parents[1]
    config = json.loads((root / "config/research_workspace.json").read_text(encoding="utf-8"))
    config["state_root"] = str(tmp_path / "state")
    w = ResearchWorkspace(root, config)
    prepared = w.prepare_research("600519", "2026-09-13")
    rid, sid = prepared["research_id"], prepared["snapshot_id"]
    _, path, pack = w.pack(rid)
    before = sha(path / "manifest.json")
    calc = Calculations(w).calculate(rid, "pe_scenarios", {"earnings":"F199", "shares":"F211", "price":"F291"}, {
        "scenarios": {"reason":"仅用于接口复算测试", "value":[dict(name=n,growth=g,multiple=m,reason="测试假设，非研究结论",valid_until="2026-12-31",invalidation=["测试变化"])
            for n,g,m in [("bear",-.1,15),("base",0,20),("bull",.1,25)]]},
        "sensitivity_growth":{"value":[-.1,0,.1],"reason":"测试"},
        "sensitivity_multiples":{"value":[15,20,25],"reason":"测试"}})
    drafts = Drafts(w)
    for n in range(1,9):
        drafts.save_section(rid,sid,n,"接口测试正文，非公司研究。数值：{{value:F199}}。{{cite:F199}}",
                            "接口测试判断，不供投资分析",["F199"],["测试条件变化"])
    drafts.save_conclusion(rid,sid,"中性观察","仅用于验证同版导出，不是黄金候选",["测试"],["测试"],["测试"],calc["artifact_id"])
    report = Reports(w).build_report(rid, ["md","html","xlsx"])
    assert report["status"] == "rendered_pending_review"
    folder = Path(report["outputs"]["md"]["path"]).parent
    saved = json.loads((folder / "report.json").read_text(encoding="utf-8"))
    assert len(saved["claims"]) == 8
    assert saved["conclusion"]["price_as_of"].startswith("2026-09-11")
    assert saved["request_metadata"]["price_time_precision"] == "date_only"
    assert saved["request_metadata"]["research_progress"]["authored_sections"] == 8
    from openpyxl import load_workbook
    book = load_workbook(report["outputs"]["xlsx"]["path"])
    assert book["核心期间表"].max_row > 100
    assert book["模型估值结果"].max_row == 4
    assert book["情景测算"].max_row == 2
    assert book["情景测算"]["B2"].value == "模型估值结果"
    assert len(saved["assumptions"]) == 6
    assert saved["audit"]["model_runs"][0]["status"] == "成功"
    assert len(saved["audit"]["model_runs"]) == 1
    assert saved["model_inputs"]["VAL.RELATIVE"]["_lineage"]["target_multiple"]
    assert "{{" not in Path(report["outputs"]["md"]["path"]).read_text(encoding="utf-8")
    assert sha(path / "manifest.json") == before
    assert w.get_task(rid)["stages"]["rendering"] == "rendered_pending_review"
    repeated = Reports(w).build_report(rid, ["md"])
    assert repeated["report_version"] == report["report_version"] + 1
    from analysis.research.authoring import Authoring
    a = Authoring(w)
    ratio = Calculations(w).calculate(rid,"ratio",{"numerator":"F199","denominator":"F199"})
    for n in range(1,9):
        a.save_section(rid,sid,n,"经营变化的研究解释。{{cite:F199}}比例{{value:"+ratio['artifact_id']+".value}}","测试判断",["F199",ratio['artifact_id']])
    a.save_conclusion(rid,sid,"中性观察","测试观点",["测试依据"],"盈利连续下降时下调评级。",calc["artifact_id"])
    v2 = Reports(w).build_report(rid,["md","html"])
    body = Path(v2['outputs']['md']['path']).read_text('utf8')
    assert body.count('下行风险与评级失效触发条件') == 1
    assert body.count('盈利连续下降时下调评级。') == 1
    assert body.count('100.00%') == 8
    assert '{{value:' not in body
    assert body.index('## 6.') < body.index('下行风险与评级失效触发条件') < body.index('## 7.')
    assert len(Drafts(w).get_draft(rid)['sections']) == 8
    a.save_section(rid,sid,1,"比例{{value:"+ratio['artifact_id']+".missing}}","测试判断",[ratio['artifact_id']])
    with pytest.raises(ResearchError,match="unknown_calculation_value"):
        Reports(w).build_report(rid,["md"])
    drafts.save_section(rid,sid,1,"未知引用 {{cite:missing}}", "测试", ["F199"], ["测试"])
    with pytest.raises(ResearchError, match="mixed_writing_contract|unknown_citation"):
        Reports(w).build_report(rid, ["md"])
