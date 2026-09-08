import json
import os
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from analysis.exports import (
    _promote_generated_pdf,
    export_report,
    render_html,
    render_markdown,
)
from analysis.models import (
    DimensionalFactRecord,
    EvidenceSpan,
    EventRecord,
    ReportCreateRequest,
    VerificationStatus,
)
from analysis.reporting import ReportBuilder


def test_markdown_html_and_xlsx_share_core_result(demo_report, tmp_path):
    markdown = render_markdown(demo_report)
    html = render_html(demo_report)
    base = f"{demo_report.conclusion.fair_value_base:.4f}"
    assert base in markdown
    assert base in html
    assert markdown.index("## 置顶结论卡") < markdown.index("## 1. 公司业务与商业模式")
    assert markdown.index("## 8. 情景假设与结论") < markdown.index("## 审计附录")
    assert markdown.count("\n## ") >= 10
    assert "输入血缘" in markdown
    assert "| 情景 | 年度 | 收入增速 |" in markdown
    assert "assumption_lineage" not in markdown
    assert "<th>年度</th>" in html
    assert "运行信息" in html
    assert "assumption_lineage" not in html

    path = export_report(demo_report, "xlsx", tmp_path)
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["置顶结论", "八步正文", "财务事实", "经营维度事实", "公司事件", "情景假设", "情景测算", "估值模型", "来源与审计", "方法与版本", "八步数据覆盖", "检查"]
    assert workbook["置顶结论"]["C10"].value == demo_report.conclusion.fair_value_base
    formulas = [cell.value for row in workbook["情景测算"].iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert len(formulas) >= 20
    assert any("'情景假设'!" in formula for formula in formulas)
    assert not any("#REF!" in formula for formula in formulas)
    assert workbook["情景假设"]["C2"].font.color.rgb.endswith("0000FF")
    assert [workbook["检查"].cell(1, column).value for column in range(2, 5)] == [
        "实际差异",
        "期望差异",
        "偏差",
    ]


def test_exports_share_frozen_research_coverage_projection(demo_request, tmp_path):
    snapshot_id = "coverage-snapshot-export-001"
    demo_request.research_coverage_snapshot_id = snapshot_id
    demo_request.research_coverage = {
        "overall_status": "pending",
        "analysis_scope": {"FIN": {"years": ["Y2021", "Y2022"]}},
        "questions": [
            {
                "step_id": "ES02",
                "question_id": "ES02.Q01",
                "state": "pending",
                "applicability": "true",
                "required_requirement_ids": ["REQ.ES02.Q01.001", "REQ.ES02.Q01.002"],
                "ready_requirement_ids": ["REQ.ES02.Q01.001"],
                "missing_requirement_ids": ["REQ.ES02.Q01.002"],
                "optional_missing_ids": ["REQ.ES02.Q01.003"],
                "next_paths": ["GAP02:补充历史季度"],
                "method_status": "skeleton",
                "assumption_status": "not_required",
                "analysis_status": "not_started",
            }
        ],
    }
    report = ReportBuilder().build(demo_request)

    markdown = render_markdown(report)
    html = render_html(report)
    path = export_report(report, "xlsx", tmp_path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["八步数据覆盖"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}

    assert snapshot_id in markdown
    assert snapshot_id in html
    assert "REQ.ES02.Q01.002" in markdown
    assert "REQ.ES02.Q01.002" in html
    assert row["覆盖快照ID"] == snapshot_id
    assert row["问题"] == "ES02.Q01"
    assert json.loads(row["待补必需项"]) == ["REQ.ES02.Q01.002"]
    assert row["方法状态"] == "skeleton"


def test_xlsx_writes_complete_dimensional_fact_lineage(demo_request, tmp_path):
    official = demo_request.sources[0]
    dimensional_fact = DimensionalFactRecord(
        dimensional_fact_id="dimension-export-001",
        ticker=demo_request.ticker,
        metric_id="segment_revenue",
        dimension_type="product",
        dimension_code="product-a",
        dimension_name="产品A",
        value=8_000_000_000.0,
        unit="CNY",
        period_start=demo_request.facts[0].period_end.replace(month=1, day=1),
        period_end=demo_request.facts[0].period_end,
        period_type="annual",
        available_at=demo_request.as_of,
        source_ids=[official.source_id],
        document_ids=["document-export-001"],
        document_page=88,
        table_title="分产品营业收入",
        row_label="产品A",
        column_label="营业收入",
        evidence_spans=[
            EvidenceSpan(
                document_id="document-export-001",
                page=88,
                text="产品A营业收入为8,000,000,000元。",
                table_title="分产品营业收入",
                row_label="产品A",
                column_label="营业收入",
                start_offset=120,
                end_offset=148,
            )
        ],
        verification_status=VerificationStatus.AUTHORITATIVE_SINGLE,
        data_snapshot_id="dimension-export-snapshot",
        metadata={"fixture": True},
    )
    demo_request.dimensional_facts = [dimensional_fact]
    report = ReportBuilder().build(demo_request)

    path = export_report(report, "xlsx", tmp_path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["经营维度事实"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}

    assert sheet.max_row == 2
    assert row["dimensional_fact_id"] == dimensional_fact.dimensional_fact_id
    assert row["value"] == dimensional_fact.value
    assert row["source_ids"] == official.source_id
    assert row["document_ids"] == "document-export-001"
    assert row["page"] == "88"
    assert row["table_title"] == "分产品营业收入"
    assert row["row_label"] == "产品A"
    assert row["column_label"] == "营业收入"
    assert row["evidence_text"] == "产品A营业收入为8,000,000,000元。"
    assert row["evidence_offsets"] == "120:148"
    assert row["status"] == VerificationStatus.AUTHORITATIVE_SINGLE.value
    assert row["data_snapshot_id"] == "dimension-export-snapshot"
    assert sheet["H2"].comment is not None
    assert official.url in sheet["H2"].comment.text


def test_xlsx_writes_complete_event_state_and_lineage(demo_request, tmp_path):
    official = demo_request.sources[0]
    event = EventRecord(
        event_id="event-export-001",
        ticker=demo_request.ticker,
        event_type="repurchase",
        event_subtype="progress",
        announced_at=demo_request.as_of,
        available_at=demo_request.as_of,
        lifecycle_state="in_progress",
        summary="已累计回购100万股",
        amount=15_000_000.0,
        shares=1_000_000.0,
        ratio=0.001,
        event_terms={"price_upper_bound": 20.0},
        affected_metrics=["shares_outstanding"],
        source_ids=[official.source_id],
        document_ids=["document-event-export"],
        evidence_spans=[
            EvidenceSpan(
                document_id="document-event-export",
                page=5,
                text="公司已累计回购股份1,000,000股。",
                start_offset=50,
                end_offset=72,
            )
        ],
        verification_status=VerificationStatus.AUTHORITATIVE_SINGLE,
        data_snapshot_id="event-export-snapshot",
    )
    demo_request.events = [event]
    report = ReportBuilder().build(demo_request)

    path = export_report(report, "xlsx", tmp_path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["公司事件"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: sheet.cell(2, index + 1).value for index, header in enumerate(headers)}

    assert sheet.max_row == 2
    assert row["event_id"] == event.event_id
    assert row["event_type"] == "repurchase"
    assert row["lifecycle_state"] == "in_progress"
    assert row["is_current"] is True
    assert row["summary"] == event.summary
    assert row["amount"] == event.amount
    assert row["shares"] == event.shares
    assert row["ratio"] == event.ratio
    assert row["page"] == "5"
    assert row["evidence_text"] == "公司已累计回购股份1,000,000股。"
    assert row["evidence_offsets"] == "50:72"
    assert row["status"] == VerificationStatus.AUTHORITATIVE_SINGLE.value
    assert row["data_snapshot_id"] == "event-export-snapshot"
    assert sheet["L2"].comment is not None
    assert official.url in sheet["L2"].comment.text

    html = render_html(report)
    markdown = render_markdown(report)
    assert "<th>关键数值</th>" in html
    assert "<th>证据定位</th>" in html
    assert "<th>证据片段</th>" not in html
    assert "| 事件ID | 类别 | 状态/当前 | 公告日期 | 摘要 | 关键数值 | 核验 | 状态链 | 证据定位 |" in markdown
    assert event.event_id in html


@pytest.mark.skipif(os.getenv("RUN_PDF_TEST") != "1", reason="PDF视觉测试由验收脚本显式开启")
def test_pdf_export_smoke(demo_report, tmp_path):
    path = export_report(demo_report, "pdf", tmp_path)
    assert path.read_bytes().startswith(b"%PDF")
    assert path.stat().st_size > 10_000


def test_locked_browser_pdf_uses_validated_copy_fallback(monkeypatch, tmp_path):
    generated = tmp_path / "rendered.pdf"
    destination = tmp_path / "report.pdf"
    generated.write_bytes(b"%PDF-1.4\nvalidated-test-pdf")
    monkeypatch.setattr("analysis.exports._is_valid_pdf", lambda path: path.exists())

    original_replace = type(generated).replace

    def replace_with_locked_source(path, target):
        if path == generated:
            raise PermissionError("simulated Edge lock")
        return original_replace(path, target)

    monkeypatch.setattr(type(generated), "replace", replace_with_locked_source)

    assert _promote_generated_pdf(generated, destination) is True
    assert destination.read_bytes() == generated.read_bytes()
    assert not destination.with_suffix(".pdf.tmp").exists()


def test_blank_report_excel_labels_missing_scenarios_without_error_formulas(tmp_path):
    report = ReportBuilder().build(
        ReportCreateRequest(
            ticker="600000",
            company_name="空白草稿",
            industry="制造业",
            as_of=datetime(2026, 9, 3, tzinfo=timezone.utc),
        )
    )
    path = export_report(report, "xlsx", tmp_path)
    workbook = load_workbook(path, data_only=False)
    formulas = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert not any("#N/A" in formula for formula in formulas)
    assert workbook["情景测算"]["B2"].value == "暂无该数据"
    assert "C2:I2" in {str(item) for item in workbook["情景测算"].merged_cells.ranges}
    assert workbook["情景测算"]["C2"].alignment.wrap_text is True
    assert workbook["情景测算"].row_dimensions[2].height == 42
