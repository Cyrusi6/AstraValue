"""独立核对已导出的黄金候选；不创建报告、不联网、不写入原库。"""
from __future__ import annotations

import argparse
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import fitz
from openpyxl import load_workbook


def sha(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream,"sha256").hexdigest()


def verify(report_dir: Path, pack: Path):
    report=json.loads((report_dir/"report.json").read_text(encoding="utf8"))
    core=json.loads((pack/"core-pack.json").read_text(encoding="utf8"))
    manifest=json.loads((pack/"manifest.json").read_text(encoding="utf8"))
    out_manifest=json.loads((report_dir/"report-manifest.json").read_text(encoding="utf8"))
    assert report["ticker"]==core["ticker"]=="600519"
    for name,expected in manifest["output_hashes"].items():
        assert sha(pack/name)==expected,("pack_hash",name)
    checked=0
    for group in ["source_inputs","auxiliary_inputs","peer_inputs"]:
        for source in manifest.get(group,[]):
            descriptors=list(source.get("files",{}).values())+source.get("manifests",[])
            if "path" in source:descriptors.append(source)
            for d in descriptors:
                assert sha(d["path"])==d["sha256"],d["path"]
                checked+=1
    for entry in out_manifest["outputs"].values():
        assert sha(entry["path"])==entry["sha256"],entry["path"]
    facts={f["fact_id"]:f for f in report["facts"]}
    sources={s["source_id"]:s for s in report["audit"]["sources"]}
    for f in facts.values():
        assert set(f["source_ids"])<=set(sources)
        assert set(f["derived_from_fact_ids"])<=set(facts)
    markdown=next(report_dir.glob("*.md")).read_text(encoding="utf8")
    html=next(report_dir.glob("*.html")).read_text(encoding="utf8")
    doc=fitz.open(next(report_dir.glob("*.pdf")))
    pdf_text="".join(p.get_text() for p in doc)
    pdf_flat="".join(pdf_text.split())
    wb=load_workbook(next(report_dir.glob("*.xlsx")),data_only=False)
    assert not [(ws.title,c.coordinate,c.value) for ws in wb for row in ws for c in row if c.data_type=="e"]
    errors={"#REF!","#DIV/0!","#VALUE!","#NAME?","#NUM!","#NULL!","#SPILL!","#CALC!"}
    assert not [(ws.title,c.coordinate,c.value) for ws in wb for row in ws for c in row if isinstance(c.value,str) and c.value in errors]
    xlsx_text="\n".join(str(c.value) for ws in wb for row in ws for c in row if c.value is not None)
    assert report["data_snapshot_id"] in markdown and report["data_snapshot_id"] in html and report["data_snapshot_id"] in xlsx_text
    numerical=0
    metrics={}
    for m in core["metrics"]:
        key=(m["metric_id"],m["period"],m["period_type"])
        assert key not in metrics or metrics[key]==m, ("conflicting_period_key",key)
        metrics[key]=m
        if m["state"]!="ready":continue
        f=m["fact"];r=facts[f["fact_id"]]
        assert Decimal(str(f["value"]))==Decimal(str(r["metadata"].get("decimal_value",r["value"])))
        assert r["period_type"]==f["period_type"] and r["unit"]==f["unit"]
        numerical+=1
    rows=[row for t in report["sections"][1]["tables"] for row in t["rows"]]
    selected_display=0
    for row in rows:
        for period,value in row.items():
            if period in {"指标","口径"} or value=="待补":continue
            ref=value.rsplit("[",1)[1].rstrip("]")
            matches=[m for m in core["metrics"] if m.get("fact_ref")==ref]
            assert matches and all(m==matches[0] for m in matches)
            m=matches[0]
            expected_kind={"年度累计":"cumulative","季度单季":"single_quarter","期末存量":"instant"}[row["口径"]]
            assert m["period_type"]==expected_kind and m["period"]==period
            f=m["fact"];v=Decimal(str(f["value"]))
            expected=f"{v/Decimal('100000000'):.2f}亿元" if f["unit"]=="CNY" else f"{v*100:.2f}%" if f["unit"]=="ratio" else None
            if expected:assert value.startswith(expected)
            assert value in markdown and value in html and value in xlsx_text
            assert "".join(value.split()) in pdf_flat, value
            selected_display+=1
    annual_q4=[]
    for metric in ["operating_income","parent_net_profit","operating_cash_flow"]:
        year=metrics[(metric,"2025-12-31","cumulative")]["fact"]
        q4=metrics[(metric,"2025-12-31","single_quarter")]["fact"]
        parents=[facts[i] for i in q4["derived_from_fact_ids"]]
        y=next(f for f in parents if f["period_end"]=="2025-12-31")
        q3=next(f for f in parents if f["period_end"]=="2025-09-30")
        dec=lambda f:Decimal(str(f.get("metadata",{}).get("decimal_value",f["value"])))
        assert abs(dec(q4)-(dec(y)-dec(q3)))<Decimal("0.01")
        assert dec(year)==dec(y)
        annual_q4.append({"metric":metric,"annual":str(dec(year)),"q4":str(dec(q4)),"q3_cumulative":str(dec(q3))})
    for metric in ["eastmoney_pe_ttm","eastmoney_pb_mrq","eastmoney_ps_ttm"]:
        m=next(m for m in core["metrics"] if m["metric_id"]==metric and m["state"]=="ready")
        expected=f"{Decimal(m['fact']['value']):.2f}倍"
        assert expected in markdown and expected in html and expected in xlsx_text and expected in pdf_flat
        cell=next(wb["财务事实"].cell(r,3) for r in range(2,wb["财务事实"].max_row+1) if wb["财务事实"].cell(r,1).value==m["fact"]["fact_id"])
        assert "%" not in cell.number_format and "倍" in cell.number_format
    for section in report["sections"]:
        assert section["claims"] or "待补" in section["summary"]
    for claim in report["claims"]:
        assert set(claim["evidence_fact_ids"])<=set(facts)
        assert set(claim["evidence_source_ids"])<=set(sources)
        assert claim["unknowns"] and claim["counter_evidence"] and claim["invalidation_conditions"]
    for e in core["evidence"]:
        assert sha(e["original_path"])==e["original_sha256"]
    questions=report["research_coverage"]["questions"]
    assert len(questions)==54
    assert all(q["state"]=="pending" for q in questions)
    assert report["conclusion"]["evidence_completeness"]==0 and report["audit"]["missing_items"]
    assert report["conclusion"]["rating"]=="暂不评级"
    return {"status":"passed","report_id":report["report_id"],"report_version":report["version"],
        "source_descriptors_verified":checked,"core_values_verified":numerical,"display_cells_verified":selected_display,
        "annual_q4_reconciliation":annual_q4,"pdf_pages":len(doc),"excel_error_cells":0,
        "claim_count":len(report["claims"]),"pending_questions":54,"audit_missing_count":len(report["audit"]["missing_items"]),
        "network":"read-only validation; no acquisition called","manual_acceptance":"pending"}


if __name__=="__main__":
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path,required=True);parser.add_argument("--pack",type=Path,required=True);parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args();result=verify(args.report,args.pack);args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf8");print(json.dumps(result,ensure_ascii=False))
